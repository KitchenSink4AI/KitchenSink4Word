"""Headers/footers, sections, page numbers, images, TOC."""

import shutil
import struct
import zlib
from pathlib import Path

import pytest
from docx import Document

from word_mcp.core.errors import TargetNotFound, WordMcpError
from word_mcp.core.package import DocxPackage, qn
from word_mcp.ops import furniture as f, media, read, toc

CORPUS = Path(__file__).resolve().parents[1] / "corpus"


def make_png(path: Path, w=80, h=40, rgb=(200, 30, 30)):
    """Minimal valid PNG without external deps."""
    def chunk(tag, data):
        c = tag + data
        return struct.pack(">I", len(data)) + c + struct.pack(
            ">I", zlib.crc32(c) & 0xFFFFFFFF
        )

    raw = b"".join(
        b"\x00" + bytes(rgb) * w for _ in range(h)
    )
    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    path.write_bytes(png)
    return path


@pytest.fixture
def long_doc(tmp_path):
    dst = tmp_path / "long_report.docx"
    shutil.copy(CORPUS / "long_report.docx", dst)
    return dst


@pytest.fixture
def tables_doc(tmp_path):
    dst = tmp_path / "report_tables.docx"
    shutil.copy(CORPUS / "report_tables.docx", dst)
    return dst


@pytest.fixture
def images_doc(tmp_path):
    dst = tmp_path / "with_images.docx"
    shutil.copy(CORPUS / "with_images.docx", dst)
    return dst


# ------------------------------------------------------------------------ TOC


def test_real_word_toc_detected_and_deletable(long_doc):
    """ch1-3 carries a genuine Word-created TOC: detect it, read its cached
    entries, delete it cleanly."""
    pkg = DocxPackage(long_doc)
    info = toc.read_toc(pkg)
    assert info["present"]
    assert "TOC" in info["instruction"]
    assert info["cached_entries"], "a real updated TOC has cached entries"
    n_tocs = len(info["tocs"])

    toc.delete_toc(pkg, which=0)
    pkg.save(do_backup=False)
    pkg2 = DocxPackage(long_doc)
    info2 = toc.read_toc(pkg2)
    assert len(info2["tocs"]) == n_tocs - 1
    xml = pkg2.raw_part("word/document.xml").decode("utf-8")
    assert xml.count('w:fldCharType="begin"') == xml.count('w:fldCharType="end"')
    Document(str(long_doc))


def test_insert_read_delete_toc(tables_doc):
    pkg = DocxPackage(tables_doc)
    assert not toc.read_toc(pkg)["present"]
    result = toc.insert_toc(pkg, at_start=True, update_on_open=True)
    pkg.save(do_backup=False)
    assert result["toc_inserted"]

    pkg2 = DocxPackage(tables_doc)
    info = toc.read_toc(pkg2)
    assert info["present"]
    assert "TOC" in info["instruction"]
    settings = pkg2.raw_part("word/settings.xml").decode("utf-8")
    assert "updateFields" in settings
    xml = pkg2.raw_part("word/document.xml").decode("utf-8")
    assert xml.count('w:fldCharType="begin"') == xml.count('w:fldCharType="end"')
    Document(str(tables_doc))

    toc.delete_toc(pkg2)
    pkg2.save(do_backup=False)
    pkg3 = DocxPackage(tables_doc)
    assert not toc.read_toc(pkg3)["present"]
    Document(str(tables_doc))


def test_double_toc_refused(tables_doc):
    pkg = DocxPackage(tables_doc)
    toc.insert_toc(pkg, at_start=True, update_on_open=False)
    with pytest.raises(WordMcpError):
        toc.insert_toc(pkg, at_start=True)


def test_toc_styles_injected(tables_doc):
    """tables_doc lacks TOC styles entirely; insertion must create them."""
    pkg = DocxPackage(tables_doc)
    toc.insert_toc(pkg, at_start=True, update_on_open=False)
    pkg.save(do_backup=False)
    styles = DocxPackage(tables_doc).raw_part("word/styles.xml").decode("utf-8")
    for sid in ("TOCHeading", "TOC1", "TOC2", "TOC3"):
        assert sid in styles


def test_delete_paragraphs_refuses_cutting_toc(tables_doc):
    from word_mcp.ops import text as t

    pkg = DocxPackage(tables_doc)
    toc.insert_toc(pkg, at_start=True, update_on_open=False)
    pkg.save(do_backup=False)
    pkg2 = DocxPackage(tables_doc)
    # The TOC SDT sits before paragraph 0; its inner paragraphs are not body
    # paragraphs, so instead simulate: bare-field TOC would be paragraphs.
    # Build a bare complex field across paragraphs at the end.
    from lxml import etree

    body = pkg2.body()
    sect = body.find(qn("w:sectPr"))
    p1 = etree.Element(qn("w:p"))
    r1 = etree.SubElement(p1, qn("w:r"))
    etree.SubElement(r1, qn("w:fldChar")).set(qn("w:fldCharType"), "begin")
    p2 = etree.Element(qn("w:p"))
    r2 = etree.SubElement(p2, qn("w:r"))
    etree.SubElement(r2, qn("w:fldChar")).set(qn("w:fldCharType"), "end")
    sect.addprevious(p1)
    sect.addprevious(p2)
    pkg2.mark_dirty()
    n = read.get_document_info(pkg2)["paragraphs"]
    with pytest.raises(WordMcpError):
        t.delete_paragraphs(pkg2, n - 2)  # would cut the field in half
    t.delete_paragraphs(pkg2, n - 2, n - 1)  # whole field: fine


# ------------------------------------------------------------ headers/footers


def test_set_header_and_footer(tables_doc):
    pkg = DocxPackage(tables_doc)
    f.set_header_footer(pkg, "header", "The Growth Plan — Chapter 4")
    f.set_header_footer(pkg, "footer", "", include_page_number=True)
    pkg.save(do_backup=False)

    pkg2 = DocxPackage(tables_doc)
    hf = f.get_headers_footers(pkg2)
    assert any(
        "Growth Plan" in h["text"] for h in hf["headers"]
    ), hf["headers"]
    assert any(x["has_page_number_field"] for x in hf["footers"])
    Document(str(tables_doc))


def test_first_page_header(tables_doc):
    pkg = DocxPackage(tables_doc)
    f.set_header_footer(pkg, "header", "First page only", ref_type="first")
    pkg.save(do_backup=False)
    pkg2 = DocxPackage(tables_doc)
    xml = pkg2.raw_part("word/document.xml").decode("utf-8")
    assert "titlePg" in xml
    sections = f.list_sections(pkg2)
    assert any(
        h["type"] == "first" for s in sections for h in s["headers"]
    )


def test_add_page_numbers_with_start(tables_doc):
    pkg = DocxPackage(tables_doc)
    f.add_page_numbers(pkg, prefix="Page ", start_at=7)
    pkg.save(do_backup=False)
    xml = DocxPackage(tables_doc).raw_part("word/document.xml").decode("utf-8")
    assert 'w:start="7"' in xml


# ------------------------------------------------------------------- sections


def test_list_sections_real_doc(long_doc):
    sections = f.list_sections(DocxPackage(long_doc))
    assert len(sections) >= 1
    assert "page_width_pt" in sections[0]


def test_set_section_properties(tables_doc):
    pkg = DocxPackage(tables_doc)
    f.set_section_properties(
        pkg,
        orientation="landscape",
        margins_pt={"left": 90, "right": 90},
    )
    pkg.save(do_backup=False)
    s = f.list_sections(DocxPackage(tables_doc))[0]
    assert s["orientation"] == "landscape"
    assert s["page_width_pt"] > s["page_height_pt"]
    assert s["margins_pt"]["left"] == 90
    Document(str(tables_doc))


def test_add_section_break(tables_doc):
    pkg = DocxPackage(tables_doc)
    n_before = len(f.list_sections(pkg))
    f.add_section_break(pkg, after_index=10, break_type="nextPage")
    pkg.save(do_backup=False)
    assert len(f.list_sections(DocxPackage(tables_doc))) == n_before + 1
    Document(str(tables_doc))


# --------------------------------------------------------------------- images


def test_add_list_resize_image(tables_doc, tmp_path):
    png = make_png(tmp_path / "test.png")
    pkg = DocxPackage(tables_doc)
    result = media.add_image(pkg, str(png), at_end=True, width_pt=200)
    pkg.save(do_backup=False)
    assert result["width_pt"] == 200

    pkg2 = DocxPackage(tables_doc)
    imgs = media.list_images(pkg2)
    assert len(imgs) == 1
    assert imgs[0]["width_pt"] == 200
    # Aspect 80x40 -> height 100.
    assert imgs[0]["height_pt"] == 100

    media.resize_image(pkg2, 0, width_pt=100)
    pkg2.save(do_backup=False)
    imgs = media.list_images(DocxPackage(tables_doc))
    assert imgs[0]["width_pt"] == 100
    assert imgs[0]["height_pt"] == 50
    Document(str(tables_doc))


def test_list_images_real_doc(images_doc):
    imgs = media.list_images(DocxPackage(images_doc))
    assert len(imgs) == 2
    assert all(i["target"] for i in imgs)


def test_replace_image_same_type_guard(images_doc, tmp_path):
    png = make_png(tmp_path / "new.png")
    pkg = DocxPackage(images_doc)
    imgs = media.list_images(pkg)
    target_ext = imgs[0]["target"].rsplit(".", 1)[1].lower()
    if target_ext == "png":
        media.replace_image(pkg, 0, str(png))
        pkg.save(do_backup=False)
        Document(str(images_doc))
    else:
        with pytest.raises(WordMcpError):
            media.replace_image(pkg, 0, str(png))


def test_unsupported_image_type(tables_doc, tmp_path):
    bad = tmp_path / "x.svg"
    bad.write_text("<svg/>")
    pkg = DocxPackage(tables_doc)
    with pytest.raises(WordMcpError):
        media.add_image(pkg, str(bad), at_end=True)
