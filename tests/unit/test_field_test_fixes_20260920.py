"""Regression tests for the 2026-09-20 field-test fix batch.

Findings credit: the author's live dissertation-consolidation production run
(report under Developer Feedback/General Feedback/Beta Test Report -
2026-09-20 - Dissertation Consolidation.md; punchlist #860-#868).
Each test reproduces the reported defect first, then asserts the fix.
"""

from __future__ import annotations

from lxml import etree

from docx import Document

import word_mcp.server as srv
from word_mcp.core.package import DocxPackage, qn
from word_mcp.ops import read as rd


# --------------------------------------------------------------- helpers


def _build(path, texts=("Alpha", "Bravo", "Charlie")):
    d = Document()
    for t in texts:
        d.add_paragraph(t)
    d.save(str(path))
    return path


def _style_el(pkg, style_id):
    for s in pkg.root("word/styles.xml").findall(qn("w:style")):
        if s.get(qn("w:styleId")) == style_id:
            return s
    return None


def _ordered_sub(parent, local, order):
    """SubElement that lands in schema position (styles need pPr before rPr)."""
    el = etree.Element(qn(f"w:{local}"))
    rank = order.index(local)
    for child in parent:
        name = etree.QName(child).localname
        if name in order and order.index(name) > rank:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


_STYLE_ORDER = [
    "name", "aliases", "basedOn", "next", "link", "autoRedefine", "hidden",
    "uiPriority", "semiHidden", "unhideWhenUsed", "qFormat", "locked",
    "personal", "personalCompose", "personalReply", "rsid", "pPr", "rPr",
]


def _stamp_normal_style(path, *, size_pt=None, line=None, after=None):
    """Put explicit formatting on the Normal STYLE (not docDefaults) --
    exactly what python-docx-generated chapter files carry."""
    pkg = DocxPackage(path)
    s = _style_el(pkg, "Normal")
    if line is not None or after is not None:
        ppr = _ordered_sub(s, "pPr", _STYLE_ORDER)
        sp = etree.SubElement(ppr, qn("w:spacing"))
        if line is not None:
            sp.set(qn("w:line"), str(line))
            sp.set(qn("w:lineRule"), "auto")
        if after is not None:
            sp.set(qn("w:after"), str(after))
    if size_pt is not None:
        rpr = _ordered_sub(s, "rPr", _STYLE_ORDER)
        for tag in ("w:sz", "w:szCs"):
            etree.SubElement(rpr, qn(tag)).set(
                qn("w:val"), str(int(size_pt * 2))
            )
    pkg.mark_dirty("word/styles.xml")
    pkg.save(do_backup=False)


def _body_paras(path):
    pkg = DocxPackage(path)
    return [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"], pkg


def _spacing_of(p):
    sp = p.find(f"{qn('w:pPr')}/{qn('w:spacing')}")
    if sp is None:
        return None
    return {
        "line": sp.get(qn("w:line")),
        "after": sp.get(qn("w:after")),
        "lineRule": sp.get(qn("w:lineRule")),
    }


def _run_sizes(p):
    out = []
    for r in p.iter(qn("w:r")):
        sz = r.find(f"{qn('w:rPr')}/{qn('w:sz')}")
        out.append(sz.get(qn("w:val")) if sz is not None else None)
    return out


# ==================================================================
# #860 CRITICAL: insert_document formatting='source' loses formatting
# defined on a SHARED STYLE NAME (one layer up from docDefaults)
# ==================================================================


def test_insert_document_bakes_shared_style_name_differences(tmp_path):
    """Source defines its body look on the Normal STYLE (python-docx's
    habit), target's Normal says something else. By-name reconciliation
    hands the target's Normal to the carried paragraphs, so the source
    values must be baked explicit or the insertion silently reformats."""
    source = _build(tmp_path / "ch5.docx", ("Ch5 body one.", "Ch5 body two."))
    _stamp_normal_style(source, size_pt=12, line=480, after=0)
    target = _build(tmp_path / "diss.docx", ("T1",))
    _stamp_normal_style(target, size_pt=10, line=240, after=240)

    out = srv.insert_document(str(target), str(source))
    dd = out["document_defaults"]
    assert dd["differ"] is True
    assert "pPr.spacing.line" in dd["differing_properties"]
    assert "rPr.sz.val" in dd["differing_properties"]
    assert dd["paragraphs_baked"] >= 2
    assert dd["runs_baked"] >= 2, "font size was never carried"
    names = [s["style_name"] for s in dd.get("styles_reconciled", [])]
    assert "Normal" in names

    paras, _pkg = _body_paras(target)
    # Target's own paragraph untouched.
    assert _spacing_of(paras[0]) is None
    assert _run_sizes(paras[0]) == [None]
    for p in paras[1:]:
        sp = _spacing_of(p)
        assert sp is not None, "inserted paragraph lost the source look"
        assert sp["line"] == "480"
        assert sp["after"] == "0"
        assert _run_sizes(p) == ["24"], "12pt source text resolved to the target"


def test_insert_document_shared_style_no_difference_bakes_nothing(tmp_path):
    """Same style definitions on both sides: nothing baked, no report."""
    source = _build(tmp_path / "src.docx", ("Body.",))
    _stamp_normal_style(source, size_pt=12, line=480, after=0)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    _stamp_normal_style(target, size_pt=12, line=480, after=0)
    out = srv.insert_document(str(target), str(source))
    assert "document_defaults" not in out
    paras, _pkg = _body_paras(target)
    assert _spacing_of(paras[-1]) is None


def test_insert_document_inherited_chain_difference_is_baked(tmp_path):
    """The difference is two links up the basedOn chain: Body -> Normal.
    Only the shared name (Normal) differs, so the chain must be resolved."""
    source = _build(tmp_path / "src.docx", ("Styled body.",))
    _stamp_normal_style(source, line=480)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    _stamp_normal_style(target, line=240)
    # A "Body Text" style in BOTH files, identical, based on Normal.
    for path in (source, target):
        pkg = DocxPackage(path)
        root = pkg.root("word/styles.xml")
        st = etree.SubElement(root, qn("w:style"))
        st.set(qn("w:type"), "paragraph")
        st.set(qn("w:styleId"), "BodyText")
        etree.SubElement(st, qn("w:name")).set(qn("w:val"), "Body Text")
        etree.SubElement(st, qn("w:basedOn")).set(qn("w:val"), "Normal")
        pkg.mark_dirty("word/styles.xml")
        pkg.save(do_backup=False)
    pkg = DocxPackage(source)
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    ppr = etree.Element(qn("w:pPr"))
    p.insert(0, ppr)
    etree.SubElement(ppr, qn("w:pStyle")).set(qn("w:val"), "BodyText")
    pkg.mark_dirty()
    pkg.save(do_backup=False)

    srv.insert_document(str(target), str(source))
    paras, _pkg = _body_paras(target)
    assert _spacing_of(paras[-1])["line"] == "480"


def test_insert_document_cloned_style_is_not_baked(tmp_path):
    """A source style with no name match is CLONED with its definition, so
    its values still resolve the same way: nothing to bake."""
    source = _build(tmp_path / "src.docx", ("Quote.",))
    pkg = DocxPackage(source)
    root = pkg.root("word/styles.xml")
    st = etree.SubElement(root, qn("w:style"))
    st.set(qn("w:type"), "paragraph")
    st.set(qn("w:styleId"), "PullQuote")
    etree.SubElement(st, qn("w:name")).set(qn("w:val"), "Pull Quote")
    sppr = etree.SubElement(st, qn("w:pPr"))
    etree.SubElement(sppr, qn("w:spacing")).set(qn("w:line"), "300")
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    ppr = etree.Element(qn("w:pPr"))
    p.insert(0, ppr)
    etree.SubElement(ppr, qn("w:pStyle")).set(qn("w:val"), "PullQuote")
    pkg.mark_dirty()
    pkg.mark_dirty("word/styles.xml")
    pkg.save(do_backup=False)

    target = _build(tmp_path / "tgt.docx", ("T1",))
    srv.insert_document(str(target), str(source))
    paras, _pkg = _body_paras(target)
    sp = _spacing_of(paras[-1])
    assert sp is None or sp["line"] != "300", "cloned style needs no baking"


# ==================================================================
# #861: define_style on an existing style must not strip attributes
# and children the call does not address (w:default, rsid, eastAsia)
# ==================================================================


def test_define_style_preserves_default_rsid_and_eastasia(tmp_path):
    """Redefining Normal is the documented read-one-define-one round trip.
    It must not drop w:default='1' (Word's document-default marker), the
    rsid, or the eastAsia font slot the call never mentioned."""
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    s = _style_el(pkg, "Normal")
    rpr = _ordered_sub(s, "rPr", _STYLE_ORDER)
    rf = etree.SubElement(rpr, qn("w:rFonts"))
    rf.set(qn("w:ascii"), "Batang")
    rf.set(qn("w:eastAsia"), "Batang")
    pkg.mark_dirty("word/styles.xml")
    pkg.save(do_backup=False)

    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="Normal", name="Normal", style_type="paragraph",
        based_on=None,
        character_formatting={"font": "Times New Roman", "size_pt": 12},
    )
    pkg.save(do_backup=False)

    s = _style_el(DocxPackage(path), "Normal")
    assert s.get(qn("w:default")) == "1", "lost the document-default marker"
    assert s.find(qn("w:rsid")) is not None, "lost the rsid"
    rf = s.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
    assert rf.get(qn("w:ascii")) == "Times New Roman"
    assert rf.get(qn("w:eastAsia")) == "Batang", "lost the eastAsia slot"
    sz = s.find(f"{qn('w:rPr')}/{qn('w:sz')}")
    assert sz.get(qn("w:val")) == "24"


def test_define_style_never_writes_self_referential_based_on(tmp_path):
    """Redefining Normal itself must not produce a basedOn on Normal,
    whether based_on is omitted (round 2: omitted leaves the parent alone)
    or passed explicitly as its own id."""
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="Normal", name="Normal",
        character_formatting={"size_pt": 12},
    )
    out = sx.define_style(
        pkg, style_id="Normal", name="Normal", based_on="Normal",
        character_formatting={"size_pt": 12},
    )
    pkg.save(do_backup=False)
    s = _style_el(DocxPackage(path), "Normal")
    assert s.find(qn("w:basedOn")) is None
    assert out.get("based_on_self_skipped") is True


def test_define_style_keeps_unaddressed_children_and_reports(tmp_path):
    """Redefining a style addressing only character formatting keeps the
    paragraph formatting, uiPriority and next-style it already had."""
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="Body", name="Body Copy", style_type="paragraph",
        paragraph_formatting={"line_spacing": 2.0, "space_after_pt": 0},
        character_formatting={"size_pt": 12},
    )
    s = _style_el(pkg, "Body")
    _ordered_sub(s, "uiPriority", _STYLE_ORDER).set(qn("w:val"), "9")
    pkg.save(do_backup=False)

    pkg = DocxPackage(path)
    out = sx.define_style(
        pkg, style_id="Body", name="Body Copy", style_type="paragraph",
        character_formatting={"bold": True},
    )
    pkg.save(do_backup=False)
    assert out["replaced"] is True
    s = _style_el(DocxPackage(path), "Body")
    assert s.find(qn("w:uiPriority")).get(qn("w:val")) == "9"
    sp = s.find(f"{qn('w:pPr')}/{qn('w:spacing')}")
    assert sp is not None and sp.get(qn("w:line")) == "480"
    assert s.find(f"{qn('w:rPr')}/{qn('w:b')}") is not None
    # the size the earlier call set is still there (not addressed now)
    assert s.find(f"{qn('w:rPr')}/{qn('w:sz')}").get(qn("w:val")) == "24"


def test_define_style_refuses_changing_an_existing_style_type(tmp_path):
    from word_mcp.ops import structure as sx
    from word_mcp.core.errors import WordMcpError
    import pytest

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    with pytest.raises(WordMcpError, match="type"):
        sx.define_style(
            pkg, style_id="Normal", name="Normal", style_type="character",
        )


# ==================================================================
# #862: com_refresh_fields must reach header and footer stories
# ==================================================================


def _word_available():
    import sys

    if sys.platform != "win32":
        return False
    try:
        import pythoncom  # noqa: F401
        import win32com.client  # noqa: F401
    except ImportError:
        return False
    import winreg

    try:
        winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application")
        return True
    except OSError:
        return False


def test_story_counts_never_sum_the_same_fields_twice():
    """Pure-Python half: two update passes over the same fields report the
    larger per-story count, not the sum."""
    from word_mcp.com import bridge

    merged = bridge._merge_counts(
        {"body": 106, "primary_header": 2}, {"body": 106, "primary_footer": 2}
    )
    assert merged == {"body": 106, "primary_header": 2, "primary_footer": 2}
    assert bridge._STORY_NAMES[7] == "primary_header"
    assert bridge._STORY_NAMES[9] == "primary_footer"


def _two_section_doc_with_stale_headers(path):
    """Two sections, each with a header field whose cached result is stale."""
    from word_mcp.ops import furniture as fu
    from word_mcp.ops import structure as sx

    d = Document()
    for i in range(40):
        d.add_paragraph(f"Section one paragraph {i} " + "lorem ipsum " * 8)
    d.add_section()
    for i in range(40):
        d.add_paragraph(f"Section two paragraph {i} " + "lorem ipsum " * 8)
    d.save(str(path))

    pkg = DocxPackage(path)
    sx.set_document_properties(pkg, title="Fresh Title")
    fu.set_header_footer(pkg, "header", "sec one", section=0)
    fu.set_header_footer(pkg, "header", "sec two", section=1)
    pkg.save(do_backup=False)

    pkg = DocxPackage(path)
    for i, part in enumerate(("word/header1.xml", "word/header2.xml")):
        root = pkg.root(part)
        p = root.find(qn("w:p"))
        for r in list(p.findall(qn("w:r"))):
            p.remove(r)
        begin = etree.SubElement(p, qn("w:r"))
        etree.SubElement(begin, qn("w:fldChar")).set(
            qn("w:fldCharType"), "begin"
        )
        instr_run = etree.SubElement(p, qn("w:r"))
        it = etree.SubElement(instr_run, qn("w:instrText"))
        it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        it.text = " DOCPROPERTY Title \\* MERGEFORMAT "
        sep = etree.SubElement(p, qn("w:r"))
        etree.SubElement(sep, qn("w:fldChar")).set(
            qn("w:fldCharType"), "separate"
        )
        cached = etree.SubElement(p, qn("w:r"))
        etree.SubElement(cached, qn("w:t")).text = f"STALE{i + 1}"
        end = etree.SubElement(p, qn("w:r"))
        etree.SubElement(end, qn("w:fldChar")).set(qn("w:fldCharType"), "end")
        pkg.mark_dirty(part)
    pkg.save(do_backup=False)


def _header_texts(path, part):
    pkg = DocxPackage(path)
    return [t.text for t in pkg.root(part).iter(qn("w:t"))]


def test_com_refresh_fields_updates_every_section_header(tmp_path):
    """Iterating doc.StoryRanges reaches only the FIRST story of each type,
    so section two's header kept its stale cache. The NextStoryRange walk
    reaches every one, and the result reports per-story counts."""
    import pytest

    if not _word_available():
        pytest.skip("Word/pywin32 not available on this machine")
    from word_mcp.com import bridge

    path = tmp_path / "sections.docx"
    _two_section_doc_with_stale_headers(path)
    assert _header_texts(path, "word/header2.xml") == ["STALE2"]

    out = bridge.refresh_fields(str(path))
    assert out["fields_refreshed"] is True
    assert "fields_by_story" in out
    assert any(
        "header" in story for story in out["fields_by_story"]
    ), out["fields_by_story"]
    assert _header_texts(path, "word/header1.xml") == ["Fresh Title"]
    assert _header_texts(path, "word/header2.xml") == ["Fresh Title"], (
        "section two's header field was never updated"
    )


# ==================================================================
# #865: first_line_indent_pt must not leave a residual w:hanging
# ==================================================================


def _ind_of(p):
    ind = p.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    if ind is None:
        return None
    return {
        "left": ind.get(qn("w:left")),
        "firstLine": ind.get(qn("w:firstLine")),
        "hanging": ind.get(qn("w:hanging")),
    }


def test_first_line_indent_zero_clears_a_hanging_indent(tmp_path):
    """A heading that inherited left=720 hanging=720 from a reference-list
    style came out left=0 hanging=720 firstLine=0: a half-inch negative
    first line on a centred heading."""
    from word_mcp.ops import text as tx

    path = _build(tmp_path / "d.docx", ("References",))
    pkg = DocxPackage(path)
    tx.set_paragraph_format(
        pkg, [0], {"indent_left_pt": 36, "first_line_indent_pt": -36}
    )
    live = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    assert _ind_of(live)["hanging"] == "720"
    tx.set_paragraph_format(
        pkg, [0], {"indent_left_pt": 0, "first_line_indent_pt": 0}
    )
    pkg.save(do_backup=False)

    ind = _ind_of(_body_paras(path)[0][0])
    assert ind["firstLine"] == "0"
    assert ind["hanging"] is None, "residual hanging indent left in place"


def test_negative_first_line_indent_clears_a_first_line_indent(tmp_path):
    """The mirror: switching to a hanging indent must drop firstLine."""
    from word_mcp.ops import text as tx

    path = _build(tmp_path / "d.docx", ("Entry",))
    pkg = DocxPackage(path)
    tx.set_paragraph_format(pkg, [0], {"first_line_indent_pt": 18})
    tx.set_paragraph_format(pkg, [0], {"first_line_indent_pt": -36})
    pkg.save(do_backup=False)
    ind = _ind_of(_body_paras(path)[0][0])
    assert ind["hanging"] == "720"
    assert ind["firstLine"] is None


def test_define_style_first_line_indent_clears_hanging(tmp_path):
    """Same writer, style edition: redefining a hanging-indent style with a
    positive first line must not keep the hanging attribute."""
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="RefEntry", name="Ref Entry",
        paragraph_formatting={"indent_left_pt": 36, "first_line_indent_pt": -36},
    )
    sx.define_style(
        pkg, style_id="RefEntry", name="Ref Entry",
        paragraph_formatting={"first_line_indent_pt": 0},
    )
    pkg.save(do_backup=False)
    s = _style_el(DocxPackage(path), "RefEntry")
    ind = s.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    assert ind.get(qn("w:firstLine")) == "0"
    assert ind.get(qn("w:hanging")) in (None, "0"), "hanging indent survived"
    assert ind.get(qn("w:left")) == "720", "the left indent was not addressed"


# ==================================================================
# #866: rPr children must be written in CT_RPr schema order
# ==================================================================


def _rpr_children(p, run_index=0):
    runs = [r for r in p.iter(qn("w:r"))]
    rpr = runs[run_index].find(qn("w:rPr"))
    return [etree.QName(c).localname for c in rpr]


def _in_schema_order(names):
    from word_mcp.ops.text import _RPR_ORDER

    ranks = [_RPR_ORDER.index(n) for n in names if n in _RPR_ORDER]
    return ranks == sorted(ranks)


def test_insert_paragraphs_writes_rpr_in_schema_order(tmp_path):
    """The blue-insertion run produced <w:u/><w:color/>, which is backwards:
    CT_RPr sequences color before u."""
    path = _build(tmp_path / "d.docx", ("Anchor",))
    srv.insert_paragraphs(
        str(path),
        [{
            "text": "Inserted blue underlined text",
            "formatting": {"underline": True, "color": "0000FF"},
        }],
        location={"paragraph": 0},
    )
    paras, _pkg = _body_paras(path)
    names = _rpr_children(paras[1])
    assert set(names) >= {"u", "color"}
    assert _in_schema_order(names), names


def test_format_text_keeps_rpr_in_schema_order(tmp_path):
    """Adding emphasis to an existing rPr appended b and i after u, which
    is also invalid: b and i come first."""
    path = _build(tmp_path / "d.docx", ("Anchor",))
    srv.insert_paragraphs(
        str(path),
        [{
            "text": "Inserted blue underlined text",
            "formatting": {"underline": True, "color": "0000FF"},
        }],
        location={"paragraph": 0},
    )
    srv.format_text(
        str(path),
        find="Inserted blue underlined text",
        formatting={"bold": True, "italic": True},
    )
    paras, _pkg = _body_paras(path)
    names = _rpr_children(paras[1])
    assert set(names) >= {"b", "i", "u", "color"}
    assert _in_schema_order(names), names


def test_format_text_repairs_an_out_of_order_rpr(tmp_path):
    """A run whose rPr an older version wrote out of order is sorted when
    the tool next touches it."""
    from word_mcp.ops import text as tx

    path = _build(tmp_path / "d.docx", ("Target text",))
    pkg = DocxPackage(path)
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    r = p.find(qn("w:r"))
    rpr = etree.Element(qn("w:rPr"))
    r.insert(0, rpr)
    etree.SubElement(rpr, qn("w:u")).set(qn("w:val"), "single")
    etree.SubElement(rpr, qn("w:color")).set(qn("w:val"), "0000FF")
    tx.format_text(pkg, find="Target text", formatting={"bold": True})
    pkg.save(do_backup=False)
    paras, _pkg = _body_paras(path)
    names = _rpr_children(paras[0])
    assert _in_schema_order(names), names


def test_define_style_rpr_is_in_schema_order(tmp_path):
    """The same writer feeds style definitions."""
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="Blue", name="Blue",
        character_formatting={
            "underline": True, "color": "0000FF", "bold": True,
            "size_pt": 12, "font": "Times New Roman",
        },
    )
    pkg.save(do_backup=False)
    s = _style_el(DocxPackage(path), "Blue")
    names = [etree.QName(c).localname for c in s.find(qn("w:rPr"))]
    assert _in_schema_order(names), names


# ==================================================================
# #867: citation_parity must not drown a real bibliography in
# organizational-author false positives
# ==================================================================


def _parity_doc(tmp_path, body, refs):
    from docx import Document as _D

    path = tmp_path / "parity.docx"
    d = _D()
    for para in body:
        d.add_paragraph(para)
    d.add_heading("References", 1)
    for r in refs:
        d.add_paragraph(r)
    d.save(str(path))
    return path


def test_organizational_authors_are_matched_not_reported_missing(tmp_path):
    """The reported false positives: 'Archives (2003)' for 'National
    Archives. (2003).', 'War (2026)' for 'U.S. Department of War. (2026)',
    'Assembly (1977a)', and a possessive read as a surname."""
    from word_mcp.ops import citecheck

    path = _parity_doc(
        tmp_path,
        [
            "The National Archives (2003) holds the cable traffic.",
            "See also the U.S. Department of War (2026, March) report, and "
            "the National Assembly of the Republic of Korea (1977a).",
            "Bandura's (1977) account of self-efficacy is the origin.",
            "The dictionary entry is unambiguous (조선말 대사전, 1992).",
        ],
        [
            "National Archives. (2003). Cable traffic of the period.",
            "U.S. Department of War. (2026, March). Posture statement.",
            "National Assembly of the Republic of Korea. (1977a, June). "
            "Plenary record.",
            "Bandura, A. (1977). Self-efficacy. Psychological Review.",
            "조선말 대사전. (1992). 평양: 사회과학출판사.",
        ],
    )
    r = citecheck.check_citation_parity(DocxPackage(path))
    assert r["missing_references"] == [], r["missing_references"]
    assert r["missing_references_unparsed"] == [], (
        r["missing_references_unparsed"]
    )
    assert r["uncited_references"] == [], r["uncited_references"]
    assert r["parity_ok"] is True


def test_reference_entries_counts_every_entry(tmp_path):
    """reference_entries reported 197 for a list of 217 because unparsed
    entries were left out of the count."""
    from word_mcp.ops import citecheck

    path = _parity_doc(
        tmp_path,
        ["A claim (Smith, 2020)."],
        [
            "Smith, J. (2020). A title. Journal.",
            "An entry with no year at all, which cannot be keyed.",
        ],
    )
    r = citecheck.check_citation_parity(DocxPackage(path))
    assert r["reference_entries"] == 2
    assert r["reference_entries_parsed"] == 1
    assert len(r["unparsed_reference_entries"]) == 1


def test_a_genuinely_missing_surname_still_reports(tmp_path):
    """The check must stay useful: a real gap is still flagged, and in the
    confident list."""
    from word_mcp.ops import citecheck

    path = _parity_doc(
        tmp_path,
        ["Framed by Bordin (1979), extended later (Bordin, 1994)."],
        ["Bordin, E. S. (1979). The generalizability of the concept."],
    )
    r = citecheck.check_citation_parity(DocxPackage(path))
    assert r["missing_references"] == ["Bordin (1994)"]
    assert r["parity_ok"] is False


def test_prose_lead_words_are_not_read_as_authors(tmp_path):
    """'As Muller (2019) shows' must key on Muller, not on 'As Muller'."""
    from word_mcp.ops import citecheck

    path = _parity_doc(
        tmp_path,
        ["As Muller (2019) shows, the claim holds."],
        ["Muller, K. (2019). Der Titel. Zeitschrift."],
    )
    r = citecheck.check_citation_parity(DocxPackage(path))
    assert r["parity_ok"] is True
    assert r["unique_cited_works"] == 1, "'As Muller' counted as its own work"


# ==================================================================
# #864: whole-paragraph character formatting by index / anchor
# ==================================================================


def _italic_state(p):
    """(run italics, paragraph-mark italic) for one paragraph."""
    ppr = p.find(qn("w:pPr"))
    runs = []
    for r in p.iter(qn("w:r")):
        if r.getparent() is ppr:
            continue
        runs.append(r.find(f"{qn('w:rPr')}/{qn('w:i')}") is not None)
    mark = ppr is not None and ppr.find(
        f"{qn('w:rPr')}/{qn('w:i')}"
    ) is not None
    return runs, mark


def _make_italic(pkg, indices):
    from word_mcp.ops import text as tx

    tx.format_paragraphs(pkg, indices, {"italic": True})


def test_format_text_range_deitalicises_whole_paragraphs(tmp_path):
    """De-italicising 19 headings whose text also occurs in running prose:
    a find string cannot target them, and the old range form took one
    paragraph at a time."""
    path = _build(
        tmp_path / "d.docx",
        (
            "The Combined Forces Command",
            "The Combined Forces Command was established in 1978.",
            "The Combined Forces Command",
        ),
    )
    pkg = DocxPackage(path)
    _make_italic(pkg, [0, 2])
    pkg.save(do_backup=False)
    paras, _pkg = _body_paras(path)
    assert _italic_state(paras[0]) == ([True], True)

    srv.format_text(
        str(path), range={"start": 0, "end": 0}, formatting={"italic": False}
    )
    srv.format_text(
        str(path), range={"start": 2, "end": 2}, formatting={"italic": False}
    )
    paras, _pkg = _body_paras(path)
    assert _italic_state(paras[0]) == ([False], False), "paragraph mark kept"
    assert _italic_state(paras[2]) == ([False], False)
    # the body sentence, which shares the heading's text, is untouched
    assert _italic_state(paras[1]) == ([False], False)


def test_format_text_multi_paragraph_range(tmp_path):
    """A range of paragraphs formats in ONE call."""
    path = _build(tmp_path / "d.docx", ("H1", "H2", "H3", "Body"))
    pkg = DocxPackage(path)
    _make_italic(pkg, [0, 1, 2])
    pkg.save(do_backup=False)

    out = srv.format_text(
        str(path), range={"start": 0, "end": 2}, formatting={"italic": False}
    )
    assert out["formatted_paragraphs"] == [0, 1, 2]
    paras, _pkg = _body_paras(path)
    for p in paras[:3]:
        assert _italic_state(p) == ([False], False)


def test_format_text_multi_paragraph_range_refuses_find(tmp_path):
    import pytest
    from word_mcp.core.errors import WordMcpError

    path = _build(tmp_path / "d.docx", ("Alpha", "Bravo"))
    with pytest.raises(WordMcpError, match="find"):
        srv.format_text(
            str(path), range={"start": 0, "end": 1}, find="Alpha",
            formatting={"bold": True},
        )


def test_apply_edits_format_op_covers_the_paragraph_mark(tmp_path):
    """The anchor-addressed route gets the same whole-paragraph semantics."""
    path = _build(tmp_path / "d.docx", ("Chapter Four", "Body text."))
    pkg = DocxPackage(path)
    _make_italic(pkg, [0])
    pkg.save(do_backup=False)

    view = srv.get_document_view(str(path))
    line = next(
        ln for ln in view["view"].splitlines() if "Chapter Four" in ln
    )
    anchor = line.split("]")[0].lstrip("[")
    srv.apply_edits(
        str(path),
        [{"op": "format", "anchor": anchor, "formatting": {"italic": False}}],
    )
    paras, _pkg = _body_paras(path)
    assert _italic_state(paras[0]) == ([False], False)


# ==================================================================
# #863: relocating a paragraph without losing its run formatting
# ==================================================================


def _anchors(path):
    """Display anchor per body paragraph, in document order."""
    view = srv.get_document_view(str(path))
    out = []
    for line in view["view"].splitlines():
        if line.startswith("[") and "]" in line:
            out.append((line.split("]")[0].lstrip("["), line.split("] ", 1)[-1]))
    return out


def _texts(path):
    paras, _pkg = _body_paras(path)
    return [rd.paragraph_text(p) for p in paras]


def test_apply_edits_move_relocates_a_reference_entry(tmp_path):
    """Re-alphabetising a reference list: the entry moves, and its italic
    journal title and hanging indent survive (delete-and-reinsert loses
    both, which is why this went out to raw lxml)."""
    from word_mcp.ops import text as tx

    path = _build(
        tmp_path / "refs.docx",
        ("Adams, A. (2001). First.", "Zulu, Z. (2003). Third.",
         "Baker, B. (2002). Second."),
    )
    pkg = DocxPackage(path)
    tx.format_paragraphs(pkg, [1], {"italic": True})
    tx.set_paragraph_format(
        pkg, [1], {"indent_left_pt": 36, "first_line_indent_pt": -36}
    )
    pkg.save(do_backup=False)

    anchors = _anchors(path)
    zulu = anchors[1][0]
    baker = anchors[2][0]
    out = srv.apply_edits(
        str(path),
        [{"op": "move", "anchor": zulu,
          "location": {"anchor": baker, "position": "after"}}],
    )
    assert out["changed"]["0"]["moved"] == 1
    assert _texts(path) == [
        "Adams, A. (2001). First.",
        "Baker, B. (2002). Second.",
        "Zulu, Z. (2003). Third.",
    ]
    paras, _pkg = _body_paras(path)
    moved = paras[2]
    assert _italic_state(moved) == ([True], True), "run formatting was lost"
    ind = moved.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    assert ind.get(qn("w:hanging")) == "720", "the hanging indent was lost"


def test_apply_edits_move_block_of_paragraphs_before_a_target(tmp_path):
    path = _build(tmp_path / "d.docx", ("A", "B", "C", "D"))
    anchors = _anchors(path)
    out = srv.apply_edits(
        str(path),
        [{"op": "move", "anchors": [anchors[2][0], anchors[3][0]],
          "location": {"anchor": anchors[0][0], "position": "before"}}],
    )
    assert out["changed"]["0"]["moved"] == 2
    assert _texts(path) == ["C", "D", "A", "B"]


def test_apply_edits_move_to_the_document_start_and_end(tmp_path):
    path = _build(tmp_path / "d.docx", ("A", "B", "C"))
    anchors = _anchors(path)
    srv.apply_edits(
        str(path),
        [{"op": "move", "anchor": anchors[2][0], "location": {"paragraph": 0, "position": "start"}}],
    )
    assert _texts(path) == ["C", "A", "B"]
    anchors = _anchors(path)
    srv.apply_edits(
        str(path),
        [{"op": "move", "anchor": anchors[0][0], "location": {"paragraph": 0, "position": "end"}}],
    )
    assert _texts(path) == ["A", "B", "C"]


def test_apply_edits_move_refuses_its_own_destination(tmp_path):
    import pytest
    from word_mcp.core.errors import WordMcpError

    path = _build(tmp_path / "d.docx", ("A", "B"))
    anchors = _anchors(path)
    with pytest.raises(WordMcpError):
        srv.apply_edits(
            str(path),
            [{"op": "move", "anchor": anchors[0][0],
              "location": {"anchor": anchors[0][0], "position": "after"}}],
        )
    assert _texts(path) == ["A", "B"]


def test_apply_edits_move_is_serialized_as_one_element(tmp_path):
    """The moved w:p must be the SAME element, byte for byte, afterwards."""
    from lxml import etree as _et

    path = _build(tmp_path / "d.docx", ("A", "B", "C"))
    paras, _pkg = _body_paras(path)
    before = _et.tostring(paras[0])
    anchors = _anchors(path)
    srv.apply_edits(
        str(path),
        [{"op": "move", "anchor": anchors[0][0], "location": {"paragraph": 0, "position": "end"}}],
    )
    paras, _pkg = _body_paras(path)
    assert _et.tostring(paras[-1]) == before


# ==================================================================
# #868: discovery -- assembly recipes, outline candidates, per-field
# TOC caches
# ==================================================================


def test_workflows_cover_multi_file_assembly():
    from word_mcp.ops import workflows as wf

    tasks = {t["task"] for t in wf.get_workflows()["tasks"]}
    assert {
        "merge-chapters",
        "build-lists-without-heading-styles",
        "merge-reference-lists",
    } <= tasks
    out = wf.get_workflows("merge-chapters")
    tools = [s["tool"] for s in out["steps"]]
    assert tools[0] == "copy_document"
    assert "insert_document" in tools
    assert "com_refresh_fields" in tools
    for step in out["steps"]:
        assert step["why"]
    assert any("back to front" in n for n in out["notes"])


def test_toc_fields_report_their_own_cached_entries(tmp_path):
    """Two TOC-family fields in one body (no content control): each must
    report ITS OWN cache, not every TOC-styled paragraph in the file."""
    from word_mcp.ops import toc as tc

    path = _build(tmp_path / "d.docx", ("Body one.",))
    pkg = DocxPackage(path)
    body = pkg.body()

    def field(instr, entries, entry_style):
        fp = etree.SubElement(body, qn("w:p"))
        r1 = etree.SubElement(fp, qn("w:r"))
        etree.SubElement(r1, qn("w:fldChar")).set(
            qn("w:fldCharType"), "begin"
        )
        r2 = etree.SubElement(fp, qn("w:r"))
        it = etree.SubElement(r2, qn("w:instrText"))
        it.text = instr
        r3 = etree.SubElement(fp, qn("w:r"))
        etree.SubElement(r3, qn("w:fldChar")).set(
            qn("w:fldCharType"), "separate"
        )
        for text in entries:
            ep = etree.SubElement(body, qn("w:p"))
            ppr = etree.SubElement(ep, qn("w:pPr"))
            etree.SubElement(ppr, qn("w:pStyle")).set(
                qn("w:val"), entry_style
            )
            run = etree.SubElement(ep, qn("w:r"))
            etree.SubElement(run, qn("w:t")).text = text
        lastp = etree.SubElement(body, qn("w:p"))
        r5 = etree.SubElement(lastp, qn("w:r"))
        etree.SubElement(r5, qn("w:fldChar")).set(qn("w:fldCharType"), "end")

    field(r' TOC \o "1-3" \h ', ["Chapter One\t1", "Chapter Two\t20"],
          "TOC1")
    etree.SubElement(body, qn("w:p"))  # a plain paragraph between them
    field(r' TOC \h \z \c "Table" ', ["No entries found."],
          "TableofFigures")
    pkg.mark_dirty()
    pkg.save(do_backup=False)

    out = tc.read_toc(DocxPackage(path))
    assert len(out["tocs"]) == 2
    main, captions = out["tocs"]
    assert main["kind"] == "main"
    assert [e["text"] for e in main["cached_entries"]] == [
        "Chapter One\t1", "Chapter Two\t20",
    ]
    assert captions["kind"] == "caption_list"
    assert [e["text"] for e in captions["cached_entries"]] == [
        "No entries found."
    ], "the caption list reported the main TOC's cache"
    assert "note" in out


def test_outline_heuristic_ranks_candidates(tmp_path):
    """Centered bold is the chapter title (level 1, high confidence), a
    bold flush-left line is level 2, a numbered table caption is not a
    heading at all, and a centered-not-bold title is still found."""
    from word_mcp.ops import read as rdm
    from word_mcp.ops import text as tx

    path = _build(
        tmp_path / "ch.docx",
        (
            "Chapter Four",
            "Historical Background",
            "Table 3. Alliance events by year",
            "Ordinary body prose that runs on for a while and ends.",
        ),
    )
    pkg = DocxPackage(path)
    tx.format_paragraphs(pkg, [0, 1, 2], {"bold": True})
    tx.set_paragraph_format(pkg, [0], {"alignment": "center"})
    pkg.save(do_backup=False)

    outline = rdm.get_outline(DocxPackage(path), detect_formatted=True)
    by_text = {h["text"]: h for h in outline}
    assert by_text["Chapter Four"]["level"] == 1
    assert by_text["Chapter Four"]["confidence"] == "high"
    assert by_text["Historical Background"]["level"] == 2
    assert "Table 3. Alliance events by year" not in by_text, (
        "a numbered caption was reported as a heading"
    )
    assert "Ordinary body prose that runs on for a while and ends." not in by_text


def test_outline_heuristic_sees_style_inherited_centering(tmp_path):
    """The chapter title centered by its STYLE, not by direct pPr, came
    back at level 2 with the old direct-only check."""
    from word_mcp.ops import read as rdm
    from word_mcp.ops import structure as sx
    from word_mcp.ops import text as tx

    path = _build(tmp_path / "ch.docx", ("Chapter Five", "Body prose here."))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="ChapterTitle", name="Chapter Title",
        paragraph_formatting={"alignment": "center"},
        character_formatting={"bold": True},
    )
    tx.apply_style(pkg, [0], "ChapterTitle")
    pkg.save(do_backup=False)

    outline = rdm.get_outline(DocxPackage(path), detect_formatted=True)
    by_text = {h["text"]: h for h in outline}
    assert by_text["Chapter Five"]["level"] == 1
    assert by_text["Chapter Five"]["confidence"] == "high"


def test_outline_heuristic_finds_a_centered_unbolded_title(tmp_path):
    """Chapter 5's title was centered but not bold, so the bold-only
    heuristic missed it entirely."""
    from word_mcp.ops import read as rdm
    from word_mcp.ops import text as tx

    path = _build(tmp_path / "ch5.docx", ("Chapter Five", "Body prose here."))
    pkg = DocxPackage(path)
    tx.set_paragraph_format(pkg, [0], {"alignment": "center"})
    pkg.save(do_backup=False)

    outline = rdm.get_outline(DocxPackage(path), detect_formatted=True)
    by_text = {h["text"]: h for h in outline}
    assert "Chapter Five" in by_text
    assert by_text["Chapter Five"]["confidence"] == "low"


# ==================================================================
# ROUND 2 - adversarial review of PR #29 (2026-09-20)
# B1: define_style wrote w:hanging="0" next to a firstLine indent, and
# an explicit zero still suppresses firstLine (ECMA-376 17.3.1.12), so
# the style rendered no indent at all.
# ==================================================================


def test_define_style_first_line_indent_removes_hanging_attribute(tmp_path):
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="RefEntry", name="Ref Entry",
        paragraph_formatting={"indent_left_pt": 36, "first_line_indent_pt": -36},
    )
    sx.define_style(
        pkg, style_id="RefEntry", name="Ref Entry",
        paragraph_formatting={"first_line_indent_pt": 36},
    )
    pkg.save(do_backup=False)
    ind = _style_el(DocxPackage(path), "RefEntry").find(
        f"{qn('w:pPr')}/{qn('w:ind')}"
    )
    assert ind.get(qn("w:firstLine")) == "720"
    assert qn("w:hanging") not in ind.attrib, (
        "w:hanging suppresses firstLine even at 0; it must be REMOVED"
    )
    assert ind.get(qn("w:left")) == "720"


def test_define_style_hanging_indent_removes_first_line_attribute(tmp_path):
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="RefEntry", name="Ref Entry",
        paragraph_formatting={"first_line_indent_pt": 18},
    )
    sx.define_style(
        pkg, style_id="RefEntry", name="Ref Entry",
        paragraph_formatting={"first_line_indent_pt": -36},
    )
    pkg.save(do_backup=False)
    ind = _style_el(DocxPackage(path), "RefEntry").find(
        f"{qn('w:pPr')}/{qn('w:ind')}"
    )
    assert ind.get(qn("w:hanging")) == "720"
    assert qn("w:firstLine") not in ind.attrib


def test_define_style_indent_is_what_word_reports(tmp_path):
    """COM: Word itself must report the first-line indent the call asked
    for. The XML-only assertion passed on the broken output too."""
    import pytest

    if not _word_available():
        pytest.skip("Word/pywin32 not available on this machine")
    from word_mcp.ops import structure as sx
    from word_mcp.ops import text as tx

    path = _build(tmp_path / "fli.docx", ("A wrapping paragraph. " * 12,))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="FLI", name="First Line Indent", based_on="",
        paragraph_formatting={"first_line_indent_pt": 36},
    )
    tx.apply_style(pkg, [0], "FLI")
    pkg.save(do_backup=False)

    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("Word.Application")
    app.Visible = False
    app.DisplayAlerts = 0
    try:
        doc = app.Documents.Open(str(path.resolve()), False, True, False)
        try:
            first_line = float(doc.Paragraphs(1).FirstLineIndent)
        finally:
            doc.Close(0)
    finally:
        app.Quit()
        pythoncom.CoUninitialize()
    assert abs(first_line - 36.0) < 0.5, (
        f"Word reports FirstLineIndent={first_line}, expected 36"
    )


# ==================================================================
# M1: based_on defaulted to "Normal", so every update that did not
# mention it silently re-parented the style.
# ==================================================================


def test_define_style_update_keeps_the_existing_parent(tmp_path):
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(pkg, style_id="Quote", name="Quote")
    sx.define_style(pkg, style_id="MyQuote", name="My Quote", based_on="Quote")
    sx.define_style(
        pkg, style_id="MyQuote", name="My Quote",
        character_formatting={"bold": True},
    )
    pkg.save(do_backup=False)
    s = _style_el(DocxPackage(path), "MyQuote")
    assert s.find(qn("w:basedOn")).get(qn("w:val")) == "Quote", (
        "an update that never mentioned based_on re-parented the style"
    )


def test_define_style_create_defaults_to_normal_and_empty_clears(tmp_path):
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(pkg, style_id="Fresh", name="Fresh")
    assert _style_el(pkg, "Fresh").find(qn("w:basedOn")).get(
        qn("w:val")
    ) == "Normal"
    sx.define_style(pkg, style_id="Fresh", name="Fresh", based_on="")
    pkg.save(do_backup=False)
    assert _style_el(DocxPackage(path), "Fresh").find(qn("w:basedOn")) is None


def test_define_style_does_not_add_qformat_to_an_existing_style(tmp_path):
    """Adding qFormat to a deliberately non-quick style promotes it into
    Word's gallery (review m1)."""
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    s = _style_el(pkg, "Normal")
    q = s.find(qn("w:qFormat"))
    if q is not None:
        s.remove(q)
    _ordered_sub(s, "semiHidden", _STYLE_ORDER)
    pkg.mark_dirty("word/styles.xml")
    sx.define_style(
        pkg, style_id="Normal", name="Normal",
        character_formatting={"size_pt": 12},
    )
    pkg.save(do_backup=False)
    s = _style_el(DocxPackage(path), "Normal")
    assert s.find(qn("w:qFormat")) is None
    assert s.find(qn("w:semiHidden")) is not None


# ==================================================================
# M2: a falsy toggle could no longer clear a style property, and the
# result claimed replaced=True for a no-op.
# ==================================================================


def test_define_style_can_turn_a_toggle_off_again(tmp_path):
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="Emph", name="Emph", style_type="character",
        character_formatting={"bold": True, "italic": True},
    )
    out = sx.define_style(
        pkg, style_id="Emph", name="Emph", style_type="character",
        character_formatting={"bold": False},
    )
    pkg.save(do_backup=False)
    rpr = _style_el(DocxPackage(path), "Emph").find(qn("w:rPr"))
    b = rpr.find(qn("w:b"))
    assert b is not None and b.get(qn("w:val")) in ("0", "false", "off"), (
        "bold: false left the defined bold in place"
    )
    i = rpr.find(qn("w:i"))
    assert i is not None and qn("w:val") not in i.attrib, "italic was lost"
    assert out["replaced"] is True


def test_define_style_underline_false_writes_none(tmp_path):
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="U", name="U", style_type="character",
        character_formatting={"underline": True},
    )
    sx.define_style(
        pkg, style_id="U", name="U", style_type="character",
        character_formatting={"underline": False},
    )
    pkg.save(do_backup=False)
    u = _style_el(DocxPackage(path), "U").find(f"{qn('w:rPr')}/{qn('w:u')}")
    assert u is not None and u.get(qn("w:val")) == "none"


def test_define_style_no_op_does_not_claim_replaced(tmp_path):
    from word_mcp.ops import structure as sx

    path = _build(tmp_path / "d.docx", ("Body.",))
    pkg = DocxPackage(path)
    sx.define_style(
        pkg, style_id="Same", name="Same",
        character_formatting={"size_pt": 12},
    )
    out = sx.define_style(
        pkg, style_id="Same", name="Same",
        character_formatting={"size_pt": 12},
    )
    pkg.save(do_backup=False)
    assert out["replaced"] is False
    assert out.get("unchanged") is True


# ==================================================================
# M5 + N1: moving a section break, and moving inside a document under
# track changes
# ==================================================================


def _md5(path):
    import hashlib

    return hashlib.md5(open(path, "rb").read()).hexdigest()


def _section_break_doc(tmp_path):
    """Body: [Sec1 body, BREAK HOLDER(sectPr), Sec2 body]."""
    path = _build(tmp_path / "sect.docx", ("Section one body.", "Holder.",
                                           "Section two body."))
    pkg = DocxPackage(path)
    paras = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"]
    body_sect = pkg.body().find(qn("w:sectPr"))
    ppr = paras[1].find(qn("w:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("w:pPr"))
        paras[1].insert(0, ppr)
    ppr.append(etree.fromstring(etree.tostring(body_sect)))
    pkg.mark_dirty()
    pkg.save(do_backup=False)
    return path


def test_move_refuses_a_paragraph_carrying_a_section_break(tmp_path):
    import pytest
    from word_mcp.core.errors import WordMcpError

    path = _section_break_doc(tmp_path)
    before = _md5(path)
    anchors = _anchors(path)
    with pytest.raises(WordMcpError, match="section break"):
        srv.apply_edits(
            str(path),
            [{"op": "move", "anchor": anchors[1][0],
              "location": {"paragraph": 2, "position": "after"}}],
        )
    assert _md5(path) == before, "the file changed on a refusal"


def test_move_refuses_crossing_a_section_boundary(tmp_path):
    import pytest
    from word_mcp.core.errors import WordMcpError

    path = _section_break_doc(tmp_path)
    before = _md5(path)
    anchors = _anchors(path)
    with pytest.raises(WordMcpError, match="section boundary"):
        srv.apply_edits(
            str(path),
            [{"op": "move", "anchor": anchors[0][0],
              "location": {"paragraph": 2, "position": "after"}}],
        )
    assert _md5(path) == before


def test_move_crossing_a_section_is_possible_with_the_opt_in(tmp_path):
    path = _section_break_doc(tmp_path)
    anchors = _anchors(path)
    out = srv.apply_edits(
        str(path),
        [{"op": "move", "anchor": anchors[0][0], "allow_cross_section": True,
          "location": {"paragraph": 2, "position": "after"}}],
    )
    assert _texts(path) == ["Holder.", "Section two body.",
                            "Section one body."]
    assert any("section boundary" in w for w in out["warnings"])


def test_move_within_one_section_is_unaffected(tmp_path):
    path = _build(tmp_path / "plain.docx", ("A", "B", "C"))
    anchors = _anchors(path)
    srv.apply_edits(
        str(path),
        [{"op": "move", "anchor": anchors[2][0],
          "location": {"paragraph": 0, "position": "before"}}],
    )
    assert _texts(path) == ["C", "A", "B"]


def _turn_on_track_changes(path):
    pkg = DocxPackage(path)
    root = pkg.root("word/settings.xml")
    el = etree.Element(qn("w:trackChanges"))
    root.insert(0, el)
    pkg.mark_dirty("word/settings.xml")
    pkg.save(do_backup=False)


def test_move_warns_when_track_changes_is_on(tmp_path):
    """The file-mode surface applies directly; a document under review
    must not be told nothing (review N1)."""
    path = _build(tmp_path / "tc.docx", ("A", "B", "C"))
    _turn_on_track_changes(path)
    anchors = _anchors(path)
    out = srv.apply_edits(
        str(path),
        [{"op": "move", "anchor": anchors[2][0],
          "location": {"paragraph": 0, "position": "before"}}],
    )
    assert _texts(path) == ["C", "A", "B"]
    assert any("trackChanges" in w for w in out["warnings"]), out


def test_delete_warns_when_track_changes_is_on(tmp_path):
    path = _build(tmp_path / "tc.docx", ("A", "B", "C"))
    _turn_on_track_changes(path)
    anchors = _anchors(path)
    out = srv.apply_edits(
        str(path), [{"op": "delete", "anchor": anchors[1][0]}]
    )
    assert _texts(path) == ["A", "C"]
    assert any("trackChanges" in w for w in out["warnings"]), out


def test_no_track_changes_warning_on_an_ordinary_document(tmp_path):
    path = _build(tmp_path / "plain.docx", ("A", "B", "C"))
    anchors = _anchors(path)
    out = srv.apply_edits(
        str(path), [{"op": "delete", "anchor": anchors[1][0]}]
    )
    assert out["warnings"] == []


# ==================================================================
# m6 + m7: two reporting false notes found in review
# ==================================================================


def test_outline_heuristic_rejects_a_bold_colon_lead_in(tmp_path):
    """'Dr. Smith said the following:' in bold is a lead-in, not a
    heading (review m7)."""
    from word_mcp.ops import read as rdm
    from word_mcp.ops import text as tx

    path = _build(
        tmp_path / "d.docx",
        ("Dr. Smith said the following:", "The quoted material follows here."),
    )
    pkg = DocxPackage(path)
    tx.format_paragraphs(pkg, [0], {"bold": True})
    pkg.save(do_backup=False)
    outline = rdm.get_outline(DocxPackage(path), detect_formatted=True)
    assert [h["text"] for h in outline] == []


def test_citation_author_phrase_stops_at_a_paragraph_boundary(tmp_path):
    """The author phrase swallowed the preceding paragraph, so the
    unparsed citation was reported as 'Introduction\nBandura's (1977)'
    (review m6)."""
    from word_mcp.ops import citecheck

    path = _parity_doc(
        tmp_path,
        ["Introduction", "Bandura's (1977) account is the origin."],
        ["Smith, J. (2020). A title. Journal."],
    )
    r = citecheck.check_citation_parity(DocxPackage(path))
    flagged = r["missing_references"] + r["missing_references_unparsed"]
    assert flagged, "the missing citation must still be flagged"
    assert not any("\n" in f or "Introduction" in f for f in flagged), flagged


# ==================================================================
# B2 / B3 / M3 / M4 / m2 / m3: the adversarial #860 collision matrix.
# Every case from the review is here, including the ones that passed.
# ==================================================================


def _style(pkg, style_id, name, *, stype="paragraph", rpr=None, ppr=None,
           based_on=None, table_cell_rpr=None):
    """Define a style by hand (the tools normalise too much for a matrix)."""
    root = pkg.root("word/styles.xml")
    keep_attrs = {}
    for s in root.findall(qn("w:style")):
        if s.get(qn("w:styleId")) == style_id:
            keep_attrs = dict(s.attrib)
            root.remove(s)
    st = etree.SubElement(root, qn("w:style"))
    for k, v in keep_attrs.items():
        st.set(k, v)
    st.set(qn("w:type"), stype)
    st.set(qn("w:styleId"), style_id)
    etree.SubElement(st, qn("w:name")).set(qn("w:val"), name)
    if based_on:
        etree.SubElement(st, qn("w:basedOn")).set(qn("w:val"), based_on)
    if ppr:
        st.append(etree.fromstring(
            f'<w:pPr xmlns:w="{_W}">{ppr}</w:pPr>'.encode()
        ))
    if rpr:
        st.append(etree.fromstring(
            f'<w:rPr xmlns:w="{_W}">{rpr}</w:rPr>'.encode()
        ))
    if table_cell_rpr:
        st.append(etree.fromstring(
            f'<w:tblStylePr xmlns:w="{_W}" w:type="firstRow">'
            f"<w:rPr>{table_cell_rpr}</w:rPr></w:tblStylePr>".encode()
        ))
    pkg.mark_dirty("word/styles.xml")
    return st


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _apply_pstyle(pkg, index, style_id):
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][index]
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = etree.Element(qn("w:pPr"))
        p.insert(0, ppr)
    etree.SubElement(ppr, qn("w:pStyle")).set(qn("w:val"), style_id)
    pkg.mark_dirty()
    return p


def _rpr_of(p, run=0):
    runs = [r for r in p.iter(qn("w:r")) if r.getparent().tag != qn("w:pPr")]
    return runs[run].find(qn("w:rPr")) if runs else None


def _val(holder, tag):
    if holder is None:
        return None
    el = holder.find(qn(f"w:{tag}"))
    if el is None:
        return None
    return el.get(qn("w:val"), "1")


def test_c01_same_named_heading_differs_in_bold_italic_colour(tmp_path):
    """The ordinary collision: two chapter files both define 'heading 1'.
    Round 1 tracked six attributes, so bold, italic and colour changed
    with nothing baked and nothing reported (review B3)."""
    source = _build(tmp_path / "src.docx", ("Chapter Five",))
    pkg = DocxPackage(source)
    _style(pkg, "Heading1", "heading 1",
           rpr='<w:b/><w:i w:val="0"/><w:color w:val="1F4E79"/>'
               '<w:sz w:val="32"/>')
    _apply_pstyle(pkg, 0, "Heading1")
    pkg.save(do_backup=False)

    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Heading1", "heading 1",
           rpr='<w:b w:val="0"/><w:i/><w:color w:val="auto"/>'
               '<w:sz w:val="32"/>')
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    dd = out["document_defaults"]
    for prop in ("rPr.b.val", "rPr.i.val", "rPr.color.val"):
        assert prop in dd["differing_properties"], dd["differing_properties"]
    paras, _pkg = _body_paras(target)
    rpr = _rpr_of(paras[-1])
    assert _val(rpr, "b") == "1", "the heading lost its bold"
    assert _val(rpr, "i") == "0", "the heading gained the target's italic"
    assert _val(rpr, "color") == "1F4E79", "the heading lost its colour"


def test_c07_character_style_collision_carries_bold(tmp_path):
    source = _build(tmp_path / "src.docx", ("Emphatic text here.",))
    pkg = DocxPackage(source)
    _style(pkg, "MyEmph", "My Emphasis", stype="character",
           rpr='<w:b/><w:sz w:val="28"/>')
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    r = p.find(qn("w:r"))
    rpr = etree.Element(qn("w:rPr"))
    r.insert(0, rpr)
    etree.SubElement(rpr, qn("w:rStyle")).set(qn("w:val"), "MyEmph")
    pkg.mark_dirty()
    pkg.save(do_backup=False)

    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "MyEmph", "My Emphasis", stype="character",
           rpr='<w:b w:val="0"/><w:sz w:val="20"/>')
    pkg.save(do_backup=False)

    srv.insert_document(str(target), str(source))
    paras, _pkg = _body_paras(target)
    rpr = _rpr_of(paras[-1])
    assert _val(rpr, "b") == "1"
    assert _val(rpr, "sz") == "28"


def test_c03_empty_paragraph_keeps_its_mark_size(tmp_path):
    """An empty paragraph is rendered entirely by its mark, so a spacer
    took the target's line height (review M4)."""
    source = _build(tmp_path / "src.docx", ("Body.", ""))
    _stamp_normal_style(source, size_pt=24)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    _stamp_normal_style(target, size_pt=8)

    out = srv.insert_document(str(target), str(source))
    assert out["document_defaults"]["paragraph_marks_baked"] >= 1
    paras, _pkg = _body_paras(target)
    spacer = paras[-1]
    mark = spacer.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
    assert mark is not None, "the paragraph mark was never baked"
    assert _val(mark, "sz") == "48"


def _numbered_source(path, *, num_id="3", left="1080", hanging="360"):
    """A source whose list paragraphs get their indent from numbering.xml,
    which is where a list's geometry actually lives."""
    pkg = DocxPackage(path)
    numbering = (
        f'<w:numbering xmlns:w="{_W}">'
        f'<w:abstractNum w:abstractNumId="7"><w:nsid w:val="1A2B3C4D"/>'
        f'<w:multiLevelType w:val="hybridMultilevel"/>'
        f'<w:lvl w:ilvl="0"><w:start w:val="1"/>'
        f'<w:numFmt w:val="bullet"/><w:lvlText w:val="-"/>'
        f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{left}" '
        f'w:hanging="{hanging}"/></w:pPr></w:lvl></w:abstractNum>'
        f'<w:num w:numId="{num_id}"><w:abstractNumId w:val="7"/></w:num>'
        f"</w:numbering>"
    )
    pkg.set_raw_part(
        "word/numbering.xml",
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        + numbering.encode(),
    )
    for i in (0, 1):
        p = [el for k, _j, el in rd.body_items(pkg) if k == "paragraph"][i]
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = etree.Element(qn("w:pPr"))
            p.insert(0, ppr)
        numpr = etree.SubElement(ppr, qn("w:numPr"))
        etree.SubElement(numpr, qn("w:ilvl")).set(qn("w:val"), "0")
        etree.SubElement(numpr, qn("w:numId")).set(qn("w:val"), num_id)
    pkg.mark_dirty()
    pkg.save(do_backup=False)


def test_c15_numbered_paragraphs_keep_their_list_indent(tmp_path):
    """B2 (regression): the baker wrote the style chain's ind onto list
    paragraphs, overriding the numbering level and hanging the bullet at
    -360 twips, out in the left margin."""
    source = _build(tmp_path / "src.docx", ("List item one", "List item two"))
    _numbered_source(source)
    pkg = DocxPackage(source)
    s = _style_el(pkg, "Normal")
    ppr = _ordered_sub(s, "pPr", _STYLE_ORDER)
    etree.SubElement(ppr, qn("w:ind")).set(qn("w:left"), "0")
    pkg.mark_dirty("word/styles.xml")
    pkg.save(do_backup=False)

    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    s = _style_el(pkg, "Normal")
    ppr = _ordered_sub(s, "pPr", _STYLE_ORDER)
    etree.SubElement(ppr, qn("w:ind")).set(qn("w:left"), "720")
    pkg.mark_dirty("word/styles.xml")
    pkg.save(do_backup=False)

    srv.insert_document(str(target), str(source))
    paras, _pkg = _body_paras(target)
    for p in paras[1:]:
        assert p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None
        ind = p.find(f"{qn('w:pPr')}/{qn('w:ind')}")
        assert ind is None, (
            "a direct indent was baked onto a numbered paragraph; its "
            "bullet now hangs outside the text block"
        )


def test_numbered_paragraph_still_bakes_non_indent_differences(tmp_path):
    """The numbering guard must not swallow everything else."""
    source = _build(tmp_path / "src.docx", ("List item one", "List item two"))
    _numbered_source(source)
    _stamp_normal_style(source, size_pt=12, line=480)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    _stamp_normal_style(target, size_pt=8, line=240)

    srv.insert_document(str(target), str(source))
    paras, _pkg = _body_paras(target)
    assert _spacing_of(paras[-1])["line"] == "480"
    assert _run_sizes(paras[-1]) == ["24"]
    assert paras[-1].find(f"{qn('w:pPr')}/{qn('w:ind')}") is None


def test_numbering_ids_are_remapped_and_the_list_is_not_merged(tmp_path):
    """The carried list gets its own numId/abstractNumId, and a shared
    w:nsid (two files from one template) does not merge the two lists."""
    source = _build(tmp_path / "src.docx", ("List item one", "List item two"))
    _numbered_source(source, num_id="3")
    target = _build(tmp_path / "tgt.docx", ("T1", "T2"))
    _numbered_source(target, num_id="3")

    srv.insert_document(str(target), str(source))
    pkg = DocxPackage(target)
    root = pkg.root("word/numbering.xml")
    num_ids = [n.get(qn("w:numId")) for n in root.findall(qn("w:num"))]
    assert len(num_ids) == len(set(num_ids)) == 2, num_ids
    abs_ids = [
        a.get(qn("w:abstractNumId")) for a in root.findall(qn("w:abstractNum"))
    ]
    assert len(abs_ids) == len(set(abs_ids)) == 2, abs_ids
    nsids = [
        n.get(qn("w:val"))
        for a in root.findall(qn("w:abstractNum"))
        for n in a.findall(qn("w:nsid"))
    ]
    assert len(nsids) == len(set(nsids)), "two lists kept one nsid"
    carried = [
        p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}").get(qn("w:val"))
        for k, _i, p in rd.body_items(pkg) if k == "paragraph"
    ]
    assert carried[0] == carried[1] != carried[2] == carried[3], carried


def _table_source(path, style_id="TableGrid", name="Table Grid"):
    from docx import Document as _D

    d = _D()
    d.add_paragraph("Before the table.")
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate(t.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"r{r}c{c}"
    d.save(str(path))
    pkg = DocxPackage(path)
    tbl = [el for k, _i, el in rd.body_items(pkg) if k == "table"][0]
    tblpr = tbl.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = etree.Element(qn("w:tblPr"))
        tbl.insert(0, tblpr)
    st = tblpr.find(qn("w:tblStyle"))
    if st is None:
        st = etree.Element(qn("w:tblStyle"))
        tblpr.insert(0, st)
    st.set(qn("w:val"), style_id)
    pkg.mark_dirty()
    pkg.save(do_backup=False)
    return pkg


def test_c14_table_style_collision_is_imported_and_reported(tmp_path):
    """A same-named table style with a different definition cannot be
    baked (its formatting is conditional by region), so it is imported
    under a new name and the carried table is re-pointed (review M3)."""
    source = tmp_path / "src.docx"
    _table_source(source)
    pkg = DocxPackage(source)
    _style(pkg, "TableGrid", "Table Grid", stype="table",
           rpr='<w:sz w:val="16"/><w:color w:val="FF0000"/>')
    pkg.save(do_backup=False)

    target = tmp_path / "tgt.docx"
    _table_source(target)
    pkg = DocxPackage(target)
    _style(pkg, "TableGrid", "Table Grid", stype="table",
           rpr='<w:sz w:val="44"/><w:color w:val="0000FF"/>')
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    imported = out["styles"]["imported_renamed"]
    assert imported and imported[0]["source_name"] == "Table Grid"
    new_id = imported[0]["style_id"]
    pkg = DocxPackage(target)
    tables = [el for k, _i, el in rd.body_items(pkg) if k == "table"]
    refs = [
        t.find(f"{qn('w:tblPr')}/{qn('w:tblStyle')}").get(qn("w:val"))
        for t in tables
    ]
    assert refs[0] == "TableGrid", "the target's own table was re-pointed"
    assert refs[1] == new_id, "the carried table kept the target's look"
    imported_def = _style_el(pkg, new_id)
    assert imported_def.find(f"{qn('w:rPr')}/{qn('w:sz')}").get(
        qn("w:val")
    ) == "16"
    assert imported_def.find(qn("w:name")).get(
        qn("w:val")
    ) not in ("Table Grid",)


def test_identical_table_style_is_not_duplicated(tmp_path):
    """Same definition on both sides: by-name matching, no import."""
    source = tmp_path / "src.docx"
    _table_source(source)
    pkg = DocxPackage(source)
    _style(pkg, "TableGrid", "Table Grid", stype="table",
           rpr='<w:sz w:val="16"/>')
    pkg.save(do_backup=False)
    target = tmp_path / "tgt.docx"
    _table_source(target)
    pkg = DocxPackage(target)
    _style(pkg, "TableGrid", "Table Grid", stype="table",
           rpr='<w:sz w:val="16"/>')
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    assert "imported_renamed" not in out["styles"]
    pkg = DocxPackage(target)
    names = [
        s.find(qn("w:name")).get(qn("w:val"))
        for s in pkg.root("word/styles.xml").findall(qn("w:style"))
        if s.get(qn("w:type")) == "table"
    ]
    assert names.count("Table Grid") == 1


def test_paragraph_styles_are_never_renamed_on_collision(tmp_path):
    """The import trick must NOT touch paragraph styles: a renamed
    'heading 1' would drop the carried headings out of the TOC."""
    source = _build(tmp_path / "src.docx", ("Chapter Five",))
    pkg = DocxPackage(source)
    _style(pkg, "Heading1", "heading 1", rpr='<w:b/><w:sz w:val="32"/>')
    _apply_pstyle(pkg, 0, "Heading1")
    pkg.save(do_backup=False)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Heading1", "heading 1", rpr='<w:i/><w:sz w:val="20"/>')
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    assert "imported_renamed" not in out["styles"]
    pkg = DocxPackage(target)
    paras = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"]
    assert paras[-1].find(f"{qn('w:pPr')}/{qn('w:pStyle')}").get(
        qn("w:val")
    ) == "Heading1"


def test_c04_same_id_different_name_reports_no_false_difference(tmp_path):
    """m2: the source style's NAME is unmatched but its ID collides, so
    it is cloned under a fresh id and its definition survives. The report
    must not list properties that do not in fact differ."""
    source = _build(tmp_path / "src.docx", ("Styled line.",))
    pkg = DocxPackage(source)
    _style(pkg, "Custom1", "Source Special", rpr='<w:b/><w:sz w:val="30"/>')
    _apply_pstyle(pkg, 0, "Custom1")
    pkg.save(do_backup=False)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Custom1", "Target Special", rpr='<w:i/><w:sz w:val="18"/>')
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    dd = out.get("document_defaults")
    if dd is not None:
        assert "rPr.sz.val" not in dd["differing_properties"], (
            "a cloned style was diffed against the target's same-id style"
        )
    pkg = DocxPackage(target)
    paras = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"]
    new_sid = paras[-1].find(f"{qn('w:pPr')}/{qn('w:pStyle')}").get(qn("w:val"))
    assert new_sid != "Custom1"
    assert _style_el(pkg, new_sid).find(
        f"{qn('w:rPr')}/{qn('w:sz')}"
    ).get(qn("w:val")) == "30"


def test_c06_theme_fonts_are_reported_not_guessed(tmp_path):
    """m3: the source uses a theme font and the target an explicit one.
    Nothing can be written explicitly, so it goes in not_baked WITH a
    reason, and the note must not claim the source look was kept."""
    source = _build(tmp_path / "src.docx", ("Themed body.",))
    pkg = DocxPackage(source)
    _style(pkg, "Normal", "Normal",
           rpr='<w:rFonts w:asciiTheme="minorHAnsi"/>')
    pkg.save(do_backup=False)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Normal", "Normal", rpr='<w:rFonts w:ascii="Courier New"/>')
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    dd = out["document_defaults"]
    assert "rPr.rFonts.ascii" in dd["not_baked"], dd
    assert dd["not_baked_reasons"]
    assert "keep the source appearance" not in dd["note"], dd["note"]


def test_note_claims_preservation_only_when_everything_was_baked(tmp_path):
    source = _build(tmp_path / "src.docx", ("Body.",))
    _stamp_normal_style(source, size_pt=12, line=480)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    _stamp_normal_style(target, size_pt=8, line=240)
    out = srv.insert_document(str(target), str(source))
    dd = out["document_defaults"]
    assert "not_baked" not in dd
    assert "keep the source appearance" in dd["note"]


def test_outline_level_difference_is_reported_but_never_baked(tmp_path):
    """Outline level is structure, not appearance: baking it would move
    carried paragraphs in or out of the merged document's TOC."""
    source = _build(tmp_path / "src.docx", ("A line.",))
    pkg = DocxPackage(source)
    _style(pkg, "Shared", "Shared Style", ppr="")
    _apply_pstyle(pkg, 0, "Shared")
    pkg.save(do_backup=False)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Shared", "Shared Style",
           ppr='<w:outlineLvl w:val="0"/>')
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    dd = out["document_defaults"]
    assert "pPr.outlineLvl.val" in dd["not_baked"]
    paras, _pkg = _body_paras(target)
    assert paras[-1].find(f"{qn('w:pPr')}/{qn('w:outlineLvl')}") is None


def test_identical_same_id_styles_bake_nothing(tmp_path):
    """Both files define Heading1 identically: the resolver must compare
    the TARGET's own chain, not its post-transplant cycle guard, or every
    property looks different and gets baked."""
    source = _build(tmp_path / "src.docx", ("Chapter Five",))
    target = _build(tmp_path / "tgt.docx", ("T1",))
    for path in (source, target):
        pkg = DocxPackage(path)
        _style(pkg, "Heading1", "heading 1",
               rpr='<w:b/><w:sz w:val="32"/>',
               ppr='<w:spacing w:before="240" w:after="60"/>')
        pkg.save(do_backup=False)
    pkg = DocxPackage(source)
    _apply_pstyle(pkg, 0, "Heading1")
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    assert "document_defaults" not in out, out.get("document_defaults")
    paras, _pkg = _body_paras(target)
    assert _rpr_of(paras[-1]) is None, "identical styles were baked over"


def test_paragraph_mark_carries_east_asian_font(tmp_path):
    """Cases c08-c11: the mark's rFonts eastAsia went Batang -> Malgun
    Gothic, which resizes every CJK spacer line."""
    source = _build(tmp_path / "src.docx", ("Body.", ""))
    pkg = DocxPackage(source)
    _style(pkg, "Normal", "Normal",
           rpr='<w:rFonts w:ascii="Times New Roman" w:eastAsia="Batang"/>')
    pkg.save(do_backup=False)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Normal", "Normal",
           rpr='<w:rFonts w:ascii="Arial" w:eastAsia="Malgun Gothic"/>')
    pkg.save(do_backup=False)

    srv.insert_document(str(target), str(source))
    paras, _pkg = _body_paras(target)
    mark = paras[-1].find(f"{qn('w:pPr')}/{qn('w:rPr')}")
    rfonts = mark.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:eastAsia")) == "Batang"
    assert rfonts.get(qn("w:ascii")) == "Times New Roman"


def test_every_tracked_property_family_survives_a_collision(tmp_path):
    """One pass over the property families round 1 did not track: each
    must either be carried explicitly or be named in not_baked."""
    src_rpr = (
        '<w:u w:val="single"/><w:highlight w:val="yellow"/><w:caps/>'
        '<w:smallCaps w:val="0"/><w:strike/><w:vertAlign w:val="superscript"/>'
        '<w:position w:val="6"/><w:spacing w:val="20"/><w:w w:val="90"/>'
    )
    tgt_rpr = (
        '<w:u w:val="none"/><w:highlight w:val="none"/>'
        '<w:caps w:val="0"/><w:smallCaps/><w:strike w:val="0"/>'
        '<w:vertAlign w:val="baseline"/><w:position w:val="0"/>'
        '<w:spacing w:val="0"/><w:w w:val="100"/>'
    )
    src_ppr = '<w:keepNext/><w:contextualSpacing/><w:jc w:val="center"/>'
    tgt_ppr = (
        '<w:keepNext w:val="0"/><w:contextualSpacing w:val="0"/>'
        '<w:jc w:val="both"/>'
    )
    source = _build(tmp_path / "src.docx", ("Styled body text.",))
    pkg = DocxPackage(source)
    _style(pkg, "Shared", "Shared Style", rpr=src_rpr, ppr=src_ppr)
    _apply_pstyle(pkg, 0, "Shared")
    pkg.save(do_backup=False)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Shared", "Shared Style", rpr=tgt_rpr, ppr=tgt_ppr)
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    dd = out["document_defaults"]
    paras, _pkg = _body_paras(target)
    p = paras[-1]
    rpr = _rpr_of(p)
    assert _val(rpr, "u") == "single"
    assert _val(rpr, "highlight") == "yellow"
    assert _val(rpr, "caps") == "1"
    assert _val(rpr, "smallCaps") == "0"
    assert _val(rpr, "strike") == "1"
    assert _val(rpr, "vertAlign") == "superscript"
    assert _val(rpr, "position") == "6"
    assert _val(rpr, "w") == "90"
    ppr = p.find(qn("w:pPr"))
    assert _val(ppr, "keepNext") == "1"
    assert _val(ppr, "contextualSpacing") == "1"
    assert _val(ppr, "jc") == "center"
    for prop in ("rPr.u.val", "rPr.caps.val", "pPr.keepNext.val"):
        assert prop in dd["differing_properties"]


# ==================================================================
# ROUND 3 - M6: theme COLOURS across differing themes. Two files can
# carry byte-identical <w:color w:val=".." w:themeColor="accent1"/> and
# render different colours, because theme1.xml differs and the theme
# part never travels.
# ==================================================================

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _set_theme(path, **slots):
    """Patch the template's own theme1.xml colour slots. Patching rather
    than replacing keeps the fontScheme and fmtScheme Word requires."""
    pkg = DocxPackage(path)
    scheme = pkg.root("word/theme/theme1.xml").find(
        f"{{{_A_NS}}}themeElements/{{{_A_NS}}}clrScheme"
    )
    for slot, hexval in slots.items():
        el = scheme.find(f"{{{_A_NS}}}{slot}")
        if el is None:
            el = etree.SubElement(scheme, f"{{{_A_NS}}}{slot}")
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{_A_NS}}}srgbClr").set("val", hexval)
    pkg.mark_dirty("word/theme/theme1.xml")
    pkg.save(do_backup=False)


_BASE_SLOTS = {
    "dk1": "000000", "lt1": "FFFFFF", "dk2": "44546A", "lt2": "E7E6E6",
    "accent1": "4472C4", "accent2": "ED7D31", "accent3": "A5A5A5",
    "accent4": "FFC000", "accent5": "5B9BD5", "accent6": "70AD47",
    "hlink": "0563C1", "folHlink": "954F72",
}


def _themed_pair(tmp_path, *, src_accent, tgt_accent, style_rpr,
                 texts=("Themed body.",)):
    source = _build(tmp_path / "src.docx", texts)
    pkg = DocxPackage(source)
    _style(pkg, "Normal", "Normal", rpr=style_rpr)
    pkg.save(do_backup=False)
    _set_theme(source, **{**_BASE_SLOTS, "accent1": src_accent})
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Normal", "Normal", rpr=style_rpr)
    pkg.save(do_backup=False)
    _set_theme(target, **{**_BASE_SLOTS, "accent1": tgt_accent})
    return source, target


def test_m6_theme_color_from_a_shared_style_is_resolved(tmp_path):
    """The verifier's case: identical Normal rPr in both files, source
    accent1 red, target accent1 green. Round 2 reported nothing at all."""
    source, target = _themed_pair(
        tmp_path, src_accent="C00000", tgt_accent="00B050",
        style_rpr='<w:color w:val="C00000" w:themeColor="accent1"/>',
    )
    out = srv.insert_document(str(target), str(source))
    assert "accent1" in out["theme_colors"]["differing_slots"]
    paras, _pkg = _body_paras(target)
    color = _rpr_of(paras[-1]).find(qn("w:color"))
    assert color.get(qn("w:val")) == "C00000"
    assert qn("w:themeColor") not in color.attrib, (
        "the reference can still re-resolve against the target theme"
    )


def test_m6_direct_theme_color_is_resolved(tmp_path):
    """Direct formatting takes the same route: the reference is frozen."""
    source = _build(tmp_path / "src.docx", ("Themed body.",))
    pkg = DocxPackage(source)
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    r = p.find(qn("w:r"))
    rpr = etree.Element(qn("w:rPr"))
    r.insert(0, rpr)
    c = etree.SubElement(rpr, qn("w:color"))
    c.set(qn("w:val"), "4472C4")
    c.set(qn("w:themeColor"), "accent1")
    pkg.mark_dirty()
    pkg.save(do_backup=False)
    _set_theme(source, **{**_BASE_SLOTS, "accent1": "C00000"})
    target = _build(tmp_path / "tgt.docx", ("T1",))
    _set_theme(target, **{**_BASE_SLOTS, "accent1": "00B050"})

    out = srv.insert_document(str(target), str(source))
    assert out["theme_colors"]["references_frozen"] >= 1
    paras, _pkg = _body_paras(target)
    color = _rpr_of(paras[-1]).find(qn("w:color"))
    assert color.get(qn("w:val")) == "C00000"
    assert qn("w:themeColor") not in color.attrib


def test_m6_tinted_theme_color_resolves_through_the_tint(tmp_path):
    """themeTint mixes the slot toward white; the baked value must be the
    tinted colour, not the raw slot."""
    source, target = _themed_pair(
        tmp_path, src_accent="C00000", tgt_accent="00B050",
        style_rpr='<w:color w:val="C00000" w:themeColor="accent1" '
                  'w:themeTint="99"/>',
    )
    srv.insert_document(str(target), str(source))
    paras, _pkg = _body_paras(target)
    color = _rpr_of(paras[-1]).find(qn("w:color"))
    assert color.get(qn("w:val")) == "D96666"
    assert qn("w:themeTint") not in color.attrib


def test_m6_themed_shading_and_underline_colour_are_resolved(tmp_path):
    """Every colour-bearing aspect, not just w:color: shading fill and
    underline colour carry their own theme attributes."""
    source = _build(tmp_path / "src.docx", ("Shaded body.",))
    pkg = DocxPackage(source)
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    r = p.find(qn("w:r"))
    rpr = etree.Element(qn("w:rPr"))
    r.insert(0, rpr)
    u = etree.SubElement(rpr, qn("w:u"))
    u.set(qn("w:val"), "single")
    u.set(qn("w:color"), "4472C4")
    u.set(qn("w:themeColor"), "accent1")
    shd = etree.SubElement(rpr, qn("w:shd"))
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "4472C4")
    shd.set(qn("w:themeFill"), "accent1")
    pkg.mark_dirty()
    pkg.save(do_backup=False)
    _set_theme(source, **{**_BASE_SLOTS, "accent1": "C00000"})
    target = _build(tmp_path / "tgt.docx", ("T1",))
    _set_theme(target, **{**_BASE_SLOTS, "accent1": "00B050"})

    srv.insert_document(str(target), str(source))
    paras, _pkg = _body_paras(target)
    rpr = _rpr_of(paras[-1])
    u = rpr.find(qn("w:u"))
    assert u.get(qn("w:color")) == "C00000"
    assert qn("w:themeColor") not in u.attrib
    shd = rpr.find(qn("w:shd"))
    assert shd.get(qn("w:fill")) == "C00000"
    assert qn("w:themeFill") not in shd.attrib


def test_m6_hyperlink_style_colour_travels_with_the_import(tmp_path):
    """A cloned style definition carries theme references too; they must
    be frozen or the imported style takes the target's hyperlink colour."""
    source = _build(tmp_path / "src.docx", ("Link text.",))
    pkg = DocxPackage(source)
    _style(pkg, "SrcLink", "Source Link", stype="character",
           rpr='<w:color w:val="0563C1" w:themeColor="hyperlink"/>'
               '<w:u w:val="single"/>')
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    r = p.find(qn("w:r"))
    rpr = etree.Element(qn("w:rPr"))
    r.insert(0, rpr)
    etree.SubElement(rpr, qn("w:rStyle")).set(qn("w:val"), "SrcLink")
    pkg.mark_dirty()
    pkg.save(do_backup=False)
    _set_theme(source, **{**_BASE_SLOTS, "hlink": "C00000"})
    target = _build(tmp_path / "tgt.docx", ("T1",))
    _set_theme(target, **{**_BASE_SLOTS, "hlink": "00B050"})

    out = srv.insert_document(str(target), str(source))
    assert out["theme_colors"]["references_frozen"] >= 1
    pkg = DocxPackage(target)
    color = _style_el(pkg, "SrcLink").find(f"{qn('w:rPr')}/{qn('w:color')}")
    assert color.get(qn("w:val")) == "C00000"
    assert qn("w:themeColor") not in color.attrib


def test_m6_identical_themes_bake_nothing(tmp_path):
    """No theme difference, no rewriting, no bloat: the reference stays a
    reference so the merged document still follows its own theme."""
    source, target = _themed_pair(
        tmp_path, src_accent="4472C4", tgt_accent="4472C4",
        style_rpr='<w:color w:val="4472C4" w:themeColor="accent1"/>',
    )
    pkg = DocxPackage(source)
    p = [el for k, _i, el in rd.body_items(pkg) if k == "paragraph"][0]
    r = p.find(qn("w:r"))
    rpr = etree.Element(qn("w:rPr"))
    r.insert(0, rpr)
    c = etree.SubElement(rpr, qn("w:color"))
    c.set(qn("w:val"), "4472C4")
    c.set(qn("w:themeColor"), "accent1")
    pkg.mark_dirty()
    pkg.save(do_backup=False)

    out = srv.insert_document(str(target), str(source))
    assert "theme_colors" not in out
    assert "document_defaults" not in out
    paras, _pkg = _body_paras(target)
    color = _rpr_of(paras[-1]).find(qn("w:color"))
    assert color.get(qn("w:themeColor")) == "accent1", (
        "an identical theme must leave the reference alone"
    )


# ==================================================================
# ROUND 3 - m8: an identical table style must be imported ONCE, however
# many times its source is inserted.
# ==================================================================


def _styled_table_pair(tmp_path, src_sz, tgt_sz):
    source = tmp_path / "src.docx"
    _table_source(source)
    pkg = DocxPackage(source)
    _style(pkg, "BaseTable", "Base Table", stype="table",
           rpr=f'<w:sz w:val="{src_sz}"/>')
    _style(pkg, "TableGrid", "Table Grid", stype="table",
           based_on="BaseTable", rpr=f'<w:color w:val="FF0000"/>')
    pkg.save(do_backup=False)
    target = tmp_path / "tgt.docx"
    _table_source(target)
    pkg = DocxPackage(target)
    _style(pkg, "BaseTable", "Base Table", stype="table",
           rpr=f'<w:sz w:val="{tgt_sz}"/>')
    _style(pkg, "TableGrid", "Table Grid", stype="table",
           based_on="BaseTable", rpr='<w:color w:val="0000FF"/>')
    pkg.save(do_backup=False)
    return source, target


def _table_style_names(path):
    pkg = DocxPackage(path)
    return sorted(
        s.find(qn("w:name")).get(qn("w:val"))
        for s in pkg.root("word/styles.xml").findall(qn("w:style"))
        if s.get(qn("w:type")) == "table"
    )


def test_m8_repeated_inserts_import_one_copy(tmp_path):
    source, target = _styled_table_pair(tmp_path, 16, 44)
    outs = [srv.insert_document(str(target), str(source)) for _ in range(3)]
    names = _table_style_names(target)
    assert names.count("Table Grid (imported)") == 1, names
    assert names.count("Base Table (imported)") == 1, names
    assert not any("(imported 2)" in n for n in names), names
    # the second and third inserts SAY they reused the import
    assert outs[1]["styles"]["reused_imports"], outs[1]["styles"]
    assert "imported_renamed" not in outs[1]["styles"]
    pkg = DocxPackage(target)
    refs = {
        t.find(f"{qn('w:tblPr')}/{qn('w:tblStyle')}").get(qn("w:val"))
        for k, _i, t in rd.body_items(pkg) if k == "table"
    }
    assert len(refs) == 2, refs  # the target's own, plus one imported


def test_m8_a_genuinely_different_source_still_imports_its_own(tmp_path):
    """Reuse must key on the definition, not just the name."""
    source, target = _styled_table_pair(tmp_path, 16, 44)
    srv.insert_document(str(target), str(source))
    pkg = DocxPackage(source)
    _style(pkg, "TableGrid", "Table Grid", stype="table",
           based_on="BaseTable", rpr='<w:color w:val="00FF00"/>')
    pkg.save(do_backup=False)
    srv.insert_document(str(target), str(source))
    names = _table_style_names(target)
    assert names.count("Table Grid (imported)") == 1, names
    assert names.count("Table Grid (imported 2)") == 1, names


def test_m9_outline_level_reason_states_the_toc_consequence(tmp_path):
    """"structure, not appearance" does not tell a user their chapter
    heading may move a level in the table of contents (review m9)."""
    source = _build(tmp_path / "src.docx", ("A line.",))
    pkg = DocxPackage(source)
    _style(pkg, "Shared", "Shared Style")
    _apply_pstyle(pkg, 0, "Shared")
    pkg.save(do_backup=False)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Shared", "Shared Style", ppr='<w:outlineLvl w:val="0"/>')
    pkg.save(do_backup=False)

    dd = srv.insert_document(str(target), str(source))["document_defaults"]
    reason = " ".join(dd["not_baked_reasons"])
    assert "table of contents" in reason
    assert "navigation pane" in reason
    assert "set_paragraph_format" in reason


# ==================================================================
# ROUND 3 - the properties the round-2 completeness probe found still
# silent: run border, text frame, fitText, eastAsianLayout, cnfStyle.
# ==================================================================


def _collision(tmp_path, *, src_rpr=None, tgt_rpr=None, src_ppr=None,
               tgt_ppr=None):
    source = _build(tmp_path / "src.docx", ("Body text.",))
    pkg = DocxPackage(source)
    _style(pkg, "Normal", "Normal", rpr=src_rpr, ppr=src_ppr)
    pkg.save(do_backup=False)
    target = _build(tmp_path / "tgt.docx", ("T1",))
    pkg = DocxPackage(target)
    _style(pkg, "Normal", "Normal", rpr=tgt_rpr, ppr=tgt_ppr)
    pkg.save(do_backup=False)
    return srv.insert_document(str(target), str(source)), target


def test_run_border_is_carried(tmp_path):
    """A visible box around text used to vanish with no report."""
    out, target = _collision(
        tmp_path,
        src_rpr='<w:bdr w:val="single" w:sz="8" w:space="0" w:color="FF0000"/>',
        tgt_rpr='<w:bdr w:val="none"/>',
    )
    assert "rPr.bdr.val" in out["document_defaults"]["differing_properties"]
    paras, _pkg = _body_paras(target)
    bdr = _rpr_of(paras[-1]).find(qn("w:bdr"))
    assert bdr is not None and bdr.get(qn("w:val")) == "single"
    assert bdr.get(qn("w:color")) == "FF0000"


def test_run_border_removal_is_carried_as_none(tmp_path):
    """The mirror: the source has no border and the target's style draws
    one, so the carried run must say 'none' explicitly."""
    out, target = _collision(
        tmp_path,
        src_rpr='<w:sz w:val="24"/>',
        tgt_rpr='<w:bdr w:val="single" w:sz="8" w:space="0" w:color="000000"/>',
    )
    assert "rPr.bdr.val" in out["document_defaults"]["differing_properties"]
    paras, _pkg = _body_paras(target)
    bdr = _rpr_of(paras[-1]).find(qn("w:bdr"))
    assert bdr is not None and bdr.get(qn("w:val")) == "none"


def test_text_frame_is_carried_as_a_whole_element(tmp_path):
    out, target = _collision(
        tmp_path,
        src_ppr='<w:framePr w:w="2880" w:h="1440" w:hRule="exact" '
                'w:wrap="around" w:vAnchor="text" w:hAnchor="text"/>',
        tgt_ppr='<w:jc w:val="both"/>',
    )
    assert "pPr.framePr" in out["document_defaults"]["differing_properties"]
    paras, _pkg = _body_paras(target)
    frame = paras[-1].find(f"{qn('w:pPr')}/{qn('w:framePr')}")
    assert frame is not None and frame.get(qn("w:w")) == "2880"


def test_text_frame_the_target_adds_is_reported_not_guessed(tmp_path):
    out, _target = _collision(
        tmp_path,
        src_ppr='<w:jc w:val="left"/>',
        tgt_ppr='<w:framePr w:w="2880" w:h="1440" w:hRule="exact"/>',
    )
    dd = out["document_defaults"]
    assert "pPr.framePr" in dd["not_baked"]
    assert any("not safe" in r for r in dd["not_baked_reasons"])


def test_fit_text_and_east_asian_layout_are_carried(tmp_path):
    out, target = _collision(
        tmp_path,
        src_rpr='<w:fitText w:val="1440" w:id="1"/>'
                '<w:eastAsianLayout w:id="2" w:vert="1"/>',
        tgt_rpr='<w:sz w:val="24"/>',
    )
    props = out["document_defaults"]["differing_properties"]
    assert "rPr.fitText.val" in props
    assert "rPr.eastAsianLayout.vert" in props
    paras, _pkg = _body_paras(target)
    rpr = _rpr_of(paras[-1])
    assert rpr.find(qn("w:fitText")).get(qn("w:val")) == "1440"
    assert rpr.find(qn("w:eastAsianLayout")).get(qn("w:vert")) == "1"


def test_cnf_style_is_reported_but_never_baked(tmp_path):
    out, target = _collision(
        tmp_path,
        src_ppr='<w:cnfStyle w:val="100000000000"/>',
        tgt_ppr='<w:cnfStyle w:val="000000100000"/>',
    )
    dd = out["document_defaults"]
    assert "pPr.cnfStyle.val" in dd["not_baked"]
    assert any("conditional-formatting" in r for r in dd["not_baked_reasons"])
    paras, _pkg = _body_paras(target)
    assert paras[-1].find(f"{qn('w:pPr')}/{qn('w:cnfStyle')}") is None
