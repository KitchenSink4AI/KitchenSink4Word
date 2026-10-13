"""Comment authoring and revision accept/reject on the tracked_changes.docx
fixture (126 revisions by one reviewer) and threaded_comments.docx
(33 comments, threaded).
"""

import shutil
from pathlib import Path

import pytest
from docx import Document

from word_mcp.core.package import DocxPackage
from word_mcp.ops import comments as cm, read, revisions as rv

CORPUS = Path(__file__).resolve().parents[1] / "corpus"


@pytest.fixture
def tracked(tmp_path):
    dst = tmp_path / "tracked_changes.docx"
    shutil.copy(CORPUS / "tracked_changes.docx", dst)
    return dst


@pytest.fixture
def threaded_doc(tmp_path):
    dst = tmp_path / "threaded_comments.docx"
    shutil.copy(CORPUS / "threaded_comments.docx", dst)
    return dst


@pytest.fixture
def tables_doc(tmp_path):
    dst = tmp_path / "report_tables.docx"
    shutil.copy(CORPUS / "report_tables.docx", dst)
    return dst


def visible_text(path):
    pkg = DocxPackage(path)
    return "\n".join(e["text"] for e in read.get_paragraphs(pkg))


# ------------------------------------------------------------------ revisions


def test_accept_all_preserves_visible_text(tracked):
    before = visible_text(tracked)
    pkg = DocxPackage(tracked)
    n_revs = read.revision_summary(pkg)["total"]
    result = rv.accept_revisions(pkg)
    pkg.save(do_backup=False)

    assert result["revisions_resolved"] > 0
    pkg2 = DocxPackage(tracked)
    assert read.revision_summary(pkg2)["total"] == 0
    after = visible_text(tracked)
    assert after == before, "accepting all must not change what is displayed"
    Document(str(tracked))


def test_reject_all_restores_original(tracked):
    pkg = DocxPackage(tracked)
    revs = read.get_tracked_changes(pkg)
    some_insertion = next(
        (r["text"] for r in revs if r["type"] == "insertion" and len(r["text"].strip()) > 15),
        None,
    )
    some_deletion = next(
        (r["text"] for r in revs if r["type"] == "deletion" and len(r["text"].strip()) > 15),
        None,
    )
    result = rv.reject_revisions(pkg)
    pkg.save(do_backup=False)

    assert result["revisions_resolved"] > 0
    pkg2 = DocxPackage(tracked)
    assert read.revision_summary(pkg2)["total"] == 0
    after = visible_text(tracked)
    if some_deletion:
        assert some_deletion in after, "rejected deletion must restore its text"
    if some_insertion:
        assert some_insertion not in after, "rejected insertion must vanish"
    Document(str(tracked))


def test_author_filter_no_match(tracked):
    pkg = DocxPackage(tracked)
    result = rv.accept_revisions(pkg, author="Nobody Real")
    assert result["revisions_resolved"] == 0
    assert read.revision_summary(pkg)["total"] > 0


def test_author_filter_match(tracked):
    pkg = DocxPackage(tracked)
    authors = list(read.revision_summary(pkg)["by_author"])
    result = rv.accept_revisions(pkg, author=authors[0])
    pkg.save(do_backup=False)
    assert result["revisions_resolved"] > 0
    remaining = read.revision_summary(DocxPackage(tracked))["by_author"]
    assert authors[0] not in remaining


# ------------------------------------------------------------------- comments


def test_add_comment_full_stack(tables_doc):
    """tables_doc has no comments infrastructure; everything must be created."""
    pkg = DocxPackage(tables_doc)
    paras = read.get_paragraphs(pkg)
    lp = max(paras, key=lambda p: len(p["text"]))
    anchor = " ".join(lp["text"].split()[5:9])
    result = cm.add_comment(
        pkg,
        anchor_text=anchor,
        text="This claim needs a citation.",
        author="Test Author",
    )
    pkg.save(do_backup=False)

    pkg2 = DocxPackage(tables_doc)
    for part in (
        "word/comments.xml",
        "word/commentsExtended.xml",
        "word/commentsIds.xml",
        "word/people.xml",
    ):
        assert pkg2.has_part(part), part
    got = read.get_comments(pkg2)
    assert len(got) == 1
    assert got[0]["author"] == "Test Author"
    assert got[0]["text"] == "This claim needs a citation."
    assert got[0]["anchored_text"] == anchor
    assert got[0]["resolved"] is False
    Document(str(tables_doc))


def test_reply_resolve_roundtrip(tables_doc):
    pkg = DocxPackage(tables_doc)
    paras = read.get_paragraphs(pkg)
    lp = max(paras, key=lambda p: len(p["text"]))
    anchor = " ".join(lp["text"].split()[5:9])
    r1 = cm.add_comment(
        pkg, anchor_text=anchor, text="Original point.", author="Reviewer"
    )
    cm.reply_to_comment(
        pkg, comment_id=r1["comment_id"], text="Addressed in rev 2.", author="Author"
    )
    cm.resolve_comment(pkg, comment_id=r1["comment_id"])
    pkg.save(do_backup=False)

    got = read.get_comments(DocxPackage(tables_doc))
    assert len(got) == 2
    root = next(c for c in got if c["text"] == "Original point.")
    reply = next(c for c in got if c["text"] == "Addressed in rev 2.")
    assert reply["reply_to"] == root["id"]
    assert root["resolved"] is True
    Document(str(tables_doc))


def test_delete_comment_cascades_replies(threaded_doc):
    pkg = DocxPackage(threaded_doc)
    all_comments = read.get_comments(pkg)
    threaded = [c for c in all_comments if c["reply_to"] is not None]
    if not threaded:
        pytest.skip("no threaded comments in threaded_doc")
    parent_id = threaded[0]["reply_to"]
    expected_gone = {parent_id} | {
        c["id"] for c in all_comments if c["reply_to"] == parent_id
    }
    result = cm.delete_comment(pkg, comment_id=parent_id)
    pkg.save(do_backup=False)

    remaining = {c["id"] for c in read.get_comments(DocxPackage(threaded_doc))}
    assert not (expected_gone & remaining)
    assert set(result["deleted_comments"]) >= expected_gone
    Document(str(threaded_doc))


def test_delete_leaf_comment_only(threaded_doc):
    pkg = DocxPackage(threaded_doc)
    all_comments = read.get_comments(pkg)
    reply_ids = {c["reply_to"] for c in all_comments if c["reply_to"]}
    leaves = [
        c for c in all_comments if c["reply_to"] and c["id"] not in reply_ids
    ]
    if not leaves:
        pytest.skip("no leaf replies")
    victim = leaves[0]
    n = len(all_comments)
    cm.delete_comment(pkg, comment_id=victim["id"])
    pkg.save(do_backup=False)
    remaining = read.get_comments(DocxPackage(threaded_doc))
    assert len(remaining) == n - 1
    assert all(c["id"] != victim["id"] for c in remaining)


def test_comment_on_board_doc_alongside_revisions(tracked):
    """Comment + accept revisions on the same doc: both features coexist."""
    pkg = DocxPackage(tracked)
    n_comments = len(read.get_comments(pkg))
    paras = read.get_paragraphs(pkg)
    lp = max(paras, key=lambda p: len(p["text"]))
    anchor = " ".join(lp["text"].split()[2:5])
    cm.add_comment(pkg, anchor_text=anchor, text="Post-review note.", author="NA")
    rv.accept_revisions(pkg)
    pkg.save(do_backup=False)

    pkg2 = DocxPackage(tracked)
    assert len(read.get_comments(pkg2)) == n_comments + 1
    assert read.revision_summary(pkg2)["total"] == 0
    Document(str(tracked))
