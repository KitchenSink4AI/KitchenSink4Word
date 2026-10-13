"""Features: tracked-change writing, caption lists, note conversion.
(COM compare is tested in the gate script — it needs real Word.)"""

import shutil
from pathlib import Path

import pytest
from docx import Document

from word_mcp.core.errors import WordMcpError
from word_mcp.core.package import DocxPackage
from word_mcp.ops import notes, read, revisions as rv, tables as tb, text as tx, toc

CORPUS = Path(__file__).resolve().parents[1] / "corpus"


@pytest.fixture
def tables_doc(tmp_path):
    dst = tmp_path / "report_tables.docx"
    shutil.copy(CORPUS / "report_tables.docx", dst)
    return dst


@pytest.fixture
def footnotes_doc(tmp_path):
    dst = tmp_path / "many_footnotes.docx"
    shutil.copy(CORPUS / "many_footnotes.docx", dst)
    return dst


def visible(path):
    return "\n".join(e["text"] for e in read.get_paragraphs(DocxPackage(path)))


# ------------------------------------------------------- tracked replace


def test_tracked_replace_shows_as_revision(tables_doc):
    pkg = DocxPackage(tables_doc)
    result = tx.search_and_replace(
        pkg,
        [{"find": "Growth Plan", "replace": "Growth Framework"}],
        track=True,
        author="Test Author",
    )
    pkg.save(do_backup=False)
    assert result["tracked_as"] == "Test Author"
    assert result["total"] > 0

    pkg2 = DocxPackage(tables_doc)
    summary = read.revision_summary(pkg2)
    assert summary["by_author"].get("Test Author", 0) >= result["total"]
    # Visible text shows the replacement (insertions display, deletions hide).
    text = visible(tables_doc)
    assert "Growth Framework" in text
    assert "Growth Plan" not in text
    Document(str(tables_doc))


def test_tracked_replace_reject_restores_original(tables_doc):
    before = visible(tables_doc)
    pkg = DocxPackage(tables_doc)
    tx.search_and_replace(
        pkg,
        [{"find": "Growth Plan", "replace": "WRONG TERM"}],
        track=True,
        author="Reviewer",
    )
    pkg.save(do_backup=False)

    pkg2 = DocxPackage(tables_doc)
    rv.reject_revisions(pkg2, author="Reviewer")
    pkg2.save(do_backup=False)
    assert visible(tables_doc) == before
    assert read.revision_summary(DocxPackage(tables_doc))["total"] == 0


def test_tracked_replace_accept_keeps_replacement(tables_doc):
    pkg = DocxPackage(tables_doc)
    tx.search_and_replace(
        pkg,
        [{"find": "Growth Plan", "replace": "Growth Framework"}],
        track=True,
    )
    pkg.save(do_backup=False)
    before_accept = visible(tables_doc)

    pkg2 = DocxPackage(tables_doc)
    rv.accept_revisions(pkg2)
    pkg2.save(do_backup=False)
    assert visible(tables_doc) == before_accept
    assert read.revision_summary(DocxPackage(tables_doc))["total"] == 0
    assert "Growth Framework" in visible(tables_doc)


# ----------------------------------------- tracked insert/delete paragraphs


def test_tracked_insert_and_reject(tables_doc):
    before = visible(tables_doc)
    pkg = DocxPackage(tables_doc)
    tx.insert_paragraphs(
        pkg,
        [{"text": "A tracked new paragraph."}],
        after_index=3,
        track=True,
        author="Author",
    )
    pkg.save(do_backup=False)
    assert "A tracked new paragraph." in visible(tables_doc)
    assert read.revision_summary(DocxPackage(tables_doc))["by_author"]["Author"] >= 1

    pkg2 = DocxPackage(tables_doc)
    rv.reject_revisions(pkg2, author="Author")
    pkg2.save(do_backup=False)
    assert visible(tables_doc) == before


def test_tracked_delete_and_reject(tables_doc):
    before = visible(tables_doc)
    pkg = DocxPackage(tables_doc)
    paras = read.get_paragraphs(pkg)
    victim = next(p for p in paras if len(p["text"]) > 80)
    result = tx.delete_paragraphs(
        pkg, victim["index"], track=True, author="Chair"
    )
    pkg.save(do_backup=False)
    assert result["deleted_tracked"] == 1
    # Text hidden but nothing physically removed.
    assert victim["text"] not in visible(tables_doc)
    n_paras = read.get_document_info(DocxPackage(tables_doc))["paragraphs"]
    assert n_paras == len(paras)

    pkg2 = DocxPackage(tables_doc)
    rv.reject_revisions(pkg2, author="Chair")
    pkg2.save(do_backup=False)
    assert visible(tables_doc) == before


def test_tracked_delete_accept_removes(tables_doc):
    pkg = DocxPackage(tables_doc)
    paras = read.get_paragraphs(pkg)
    victim = next(p for p in paras if len(p["text"]) > 80)
    tx.delete_paragraphs(pkg, victim["index"], track=True)
    pkg.save(do_backup=False)

    pkg2 = DocxPackage(tables_doc)
    rv.accept_revisions(pkg2)
    pkg2.save(do_backup=False)
    assert victim["text"] not in visible(tables_doc)
    assert read.revision_summary(DocxPackage(tables_doc))["total"] == 0
    Document(str(tables_doc))


# ------------------------------------------------------- tracked cell edits


def test_tracked_set_cells_roundtrip(tables_doc):
    pkg = DocxPackage(tables_doc)
    tb.create_table(pkg, [["A", "B"], ["old text", "keep"]], at_end=True)
    pkg.save(do_backup=False)

    pkg2 = DocxPackage(tables_doc)
    idx = len(read.list_tables(pkg2)) - 1
    tb.set_cells(
        pkg2,
        idx,
        [{"row": 1, "cell": 0, "text": "new text"}],
        track=True,
        author="Editor",
    )
    pkg2.save(do_backup=False)
    g = read.get_table(DocxPackage(tables_doc), idx)
    assert g["cells"][1][0]["text"] == "new text"

    pkg3 = DocxPackage(tables_doc)
    rv.reject_revisions(pkg3, author="Editor")
    pkg3.save(do_backup=False)
    g = read.get_table(DocxPackage(tables_doc), idx)
    assert g["cells"][1][0]["text"] == "old text"


# ----------------------------------------------------------- caption lists


def test_insert_caption_list(tables_doc):
    from word_mcp.ops import fields as fl

    pkg = DocxPackage(tables_doc)
    tb.create_table(pkg, [["x"]], at_end=True)
    fl.add_caption(pkg, table_index=0, label="Table", text="First table")
    toc.insert_caption_list(pkg, label="Table", at_start=True, update_on_open=False)
    pkg.save(do_backup=False)

    pkg2 = DocxPackage(tables_doc)
    info = toc.read_toc(pkg2)
    assert info["present"]
    kinds = [t["kind"] for t in info["tocs"]]
    assert "caption_list" in kinds
    xml = pkg2.raw_part("word/document.xml").decode("utf-8")
    assert '\\c "Table"' in xml
    assert "List of Tables" in xml
    Document(str(tables_doc))


def test_duplicate_caption_list_refused(tables_doc):
    pkg = DocxPackage(tables_doc)
    toc.insert_caption_list(pkg, label="Figure", at_start=True, update_on_open=False)
    with pytest.raises(WordMcpError, match="already has"):
        toc.insert_caption_list(pkg, label="Figure", at_start=True)
    # A different label is fine, and so is a main TOC.
    toc.insert_caption_list(pkg, label="Table", at_start=True, update_on_open=False)
    toc.insert_toc(pkg, at_start=True, update_on_open=False)


# ---------------------------------------------------------- note conversion


def test_convert_single_footnote_to_endnote(footnotes_doc):
    pkg = DocxPackage(footnotes_doc)
    fn_before = read.list_footnotes(pkg)
    target_text = fn_before[2]["text"]
    result = notes.convert_notes(
        pkg, "footnotes_to_endnotes", position=3
    )
    pkg.save(do_backup=False)

    pkg2 = DocxPackage(footnotes_doc)
    assert result["converted"] == 1
    assert len(read.list_footnotes(pkg2)) == len(fn_before) - 1
    ens = read.list_endnotes(pkg2)
    assert len(ens) == 1
    assert ens[0]["text"] == target_text
    v = notes.validate_notes(pkg2)
    assert v["footnotes"]["ok"] and v["endnotes"]["ok"]
    Document(str(footnotes_doc))


def test_convert_all_and_back(footnotes_doc):
    pkg = DocxPackage(footnotes_doc)
    texts_before = [n["text"] for n in read.list_footnotes(pkg)]
    result = notes.convert_notes(pkg, "footnotes_to_endnotes")
    pkg.save(do_backup=False)
    assert result["converted"] == 171

    pkg2 = DocxPackage(footnotes_doc)
    assert len(read.list_footnotes(pkg2)) == 0
    ens = read.list_endnotes(pkg2)
    assert len(ens) == 171
    assert [n["text"] for n in ens] == texts_before

    notes.convert_notes(pkg2, "endnotes_to_footnotes")
    pkg2.save(do_backup=False)
    pkg3 = DocxPackage(footnotes_doc)
    assert len(read.list_endnotes(pkg3)) == 0
    fns = read.list_footnotes(pkg3)
    assert [n["text"] for n in fns] == texts_before
    v = notes.validate_notes(pkg3)
    assert v["footnotes"]["ok"] and not v["footnotes"]["needs_cleanup"]
    Document(str(footnotes_doc))


def test_convert_bad_direction(footnotes_doc):
    pkg = DocxPackage(footnotes_doc)
    with pytest.raises(WordMcpError):
        notes.convert_notes(pkg, "sideways")
